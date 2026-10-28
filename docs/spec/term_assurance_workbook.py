"""Build the reference workbook for the term assurance specification (D22, reference case 1).

This generator reads nothing but ``docs/spec/term-assurance.md``. The only numbers typed into
the workbook are that document's reference inputs: the valuation date of section 3.3, the basis
constants and short tables of section 6 and the ten model points of section 13. Every other cell
is an Excel formula transcribing one formula of the specification, in its stated evaluation order
and computation form (powers where it says power, ratios of two curve cells, left-to-right
subtraction, sums in ascending ``t``). No engine output, number script or other document is read
and nothing is rounded. Python computes a policy's ``rem_m`` only to know how many rows to lay
out; the policy sheet recomputes it by formula and its ``rows_ok`` cell compares the two.

Sheets written to ``docs/spec/term-assurance/term-assurance.xlsx``, in the layout that section 11
of the specification fixes:

``readme``        what each sheet, block and column is, with units and the spec sections.
``mortality``     the section 6.1 parameters and the mortality table materialised for ages 18 to
                  120 and both sexes from the Gompertz-Makeham formula (the materialised-block
                  option of section 11). The one block the policy sheets share: each ``q_ann``
                  cell looks it up by INDEX/MATCH on ``(sex, age)``.
``P01`` .. ``P10``  one sheet per policy, named by ``policy_id``, with these blocks from the top:
                  the inputs of 3.1 (and the run-level ``valuation_date`` of 3.3), typed; the
                  derived constants of 3.2 and 4.3 by formula; the basis constants of section 6
                  with ``i_m``, ``v_m`` and ``e_m`` by formula; the lapse and commission tables;
                  the monthly block of 7.1 and 11.1 (a header row, rows ``t = 0 .. rem_m``, one
                  column per quantity in the 7.1 order, the row-level check cells of section 12
                  to its right); the scalars of 11.2; the check cells of section 12 (the section
                  1 scales, two helper sums, one ``<id>_ok`` per (sheet) invariant, ``rows_ok``
                  and ``ok``). Sheet-scoped defined names point at the sheet's own cells, so the
                  formulas read as the spec and a reviewer checks one sheet without leaving it.
``summary``       the 11.2 scalars per policy and the 11.3 scalar totals (plain SUMs), the check
                  cells per policy repeated by reference with ``all_ok`` (defined name
                  ``all_ok``, labelled in the cell to its left), and the 11.3 totals by ``t``.

Run from the repository root, then recalculate in Excel so the cached values exist:

    uv run --frozen python docs/spec/term_assurance_workbook.py
    uv run --frozen python docs/spec/recalc_excel.py docs/spec/term-assurance/term-assurance.xlsx
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

OUT = Path("docs/spec/term-assurance/term-assurance.xlsx")

# ----------------------------------------------------------------------------------------------
# Reference inputs: the only typed numbers in the workbook.
# ----------------------------------------------------------------------------------------------

# Section 3.3: the run-level input.
VALUATION_DATE = dt.date(2027, 1, 1)

# Section 6 constants in the order section 11 lists them: (symbol, value or formula, unit, section).
# ``i_m``, ``v_m`` and ``e_m`` are formulas of ``i`` and ``e``; ``v_m`` is the direct power.
BASIS: list[tuple[str, float | str, str, str]] = [
    ("i", 0.04, "Rate[Annual]", "6"),
    ("e", 0.03, "Rate[Annual]", "6"),
    ("i_m", "=(1+i)^(1/12)-1", "Rate[Monthly]", "6"),
    ("v_m", "=(1+i)^(-1/12)", "factor", "6"),
    ("e_m", "=(1+e)^(1/12)-1", "Rate[Monthly]", "6"),
    ("A_M", 0.0002, "Gompertz-Makeham A, sex M", "6.1"),
    ("A_F", 0.00015, "Gompertz-Makeham A, sex F", "6.1"),
    ("B_M", 0.000017, "Gompertz-Makeham B, sex M", "6.1"),
    ("B_F", 0.000011, "Gompertz-Makeham B, sex F", "6.1"),
    ("c", 1.1, "Gompertz-Makeham c (both sexes); defined name gm_c", "6.1"),
    ("maint_base", 5.00, "Money per policy, per month, time-0 prices", "6.5"),
    ("claim_exp_base", 250.00, "Money per policy (per death), time-0 prices", "6.5"),
]
# Excel reserves the names R and C, so the spec's ``c`` is the defined name ``gm_c``.
EXCEL_NAME = {"c": "gm_c"}
# Section 6.1: the parameters the materialised mortality table is built from.
MORT_PARAMS = [entry for entry in BASIS if entry[3] == "6.1"]

# Section 6.2: annual lapse rate by clamped key k = min(dur_y, 5).
LAPSE_TABLE = [0.15, 0.10, 0.08, 0.06, 0.05, 0.04]
# Section 6.6: commission rate by clamped key k = min(dur_y, 1).
COMM_TABLE = [0.20, 0.02]
# Section 6.1: the table's declared age range.
MORT_AGES = range(18, 121)


@dataclass(frozen=True)
class ModelPoint:
    """One row of the section 13 table (the section 3.1 input fields)."""

    policy_id: str
    sex: str
    birth_date: dt.date
    entry_date: dt.date
    maturity_date: dt.date
    prem_term_y: int
    sum_assured: float
    premium: float
    pols_if_init: float

    @property
    def rem_m(self) -> int:
        """Section 4.3 ``rem_m``, used only to size the policy sheet (see its ``rows_ok`` cell)."""
        term_m = months_between(self.entry_date, self.maturity_date)
        dur_m = months_between(self.entry_date, VALUATION_DATE)
        return term_m - dur_m


def months_between(d1: dt.date, d2: dt.date) -> int:
    """Section 4.2: whole calendar months from ``d1`` to ``d2``."""
    return 12 * (d2.year - d1.year) + (d2.month - d1.month) - (1 if d2.day < d1.day else 0)


def d(iso: str) -> dt.date:
    return dt.date.fromisoformat(iso)


# Section 13, exactly as printed.
MODEL_POINTS = [
    ModelPoint("P01", "M", d("1986-03-15"), d("2024-07-01"), d("2044-07-01"), 20, 250000, 52.96, 1),
    ModelPoint("P02", "F", d("1991-11-02"), d("2027-01-01"), d("2052-01-01"), 25, 400000, 51.64, 1),
    ModelPoint("P03", "M", d("1966-05-20"), d("2007-01-15"), d("2027-01-15"), 20, 150000, 39.99, 1),
    ModelPoint("P04", "F", d("1976-08-09"), d("2015-03-20"), d("2045-03-20"), 20, 300000, 68.98, 1),
    ModelPoint("P05", "M", d("1975-04-18"), d("2018-02-10"), d("2028-02-10"), 9, 200000, 46.40, 1),
    ModelPoint("P06", "F", d("2003-02-28"), d("2026-11-10"), d("2036-11-10"), 10, 100000, 10.76, 1),
    ModelPoint(
        "P07", "M", d("1980-06-01"), d("2020-02-29"), d("2040-02-29"), 20, 500000, 105.55, 2.5
    ),
    ModelPoint(
        "P08", "M", d("1958-12-31"), d("2010-12-31"), d("2030-12-31"), 10, 120000, 128.12, 1
    ),
    ModelPoint(
        "P09", "F", d("1996-10-31"), d("2026-12-15"), d("2066-12-15"), 40, 1000000, 116.62, 1
    ),
    ModelPoint("P10", "M", d("1990-05-05"), d("2022-02-01"), d("2032-02-01"), 5, 200000, 39.27, 1),
]

# ----------------------------------------------------------------------------------------------
# Layout constants.
# ----------------------------------------------------------------------------------------------

# Policy sheet, the blocks above the monthly block: the same rows on every sheet. Column A holds
# the label (and ``t`` in the monthly block), B the value, C the spec section, D the unit.
MP_INPUT_FIELDS = [
    "policy_id", "sex", "birth_date", "entry_date", "maturity_date", "prem_term_y",
    "sum_assured", "premium", "pols_if_init",
]  # fmt: skip
INPUT_FIELDS = ["valuation_date", *MP_INPUT_FIELDS]
INPUT_META: dict[str, tuple[str, str]] = {  # section, unit
    "valuation_date": ("3.3", "Date, first of a month (run-level input)"),
    "policy_id": ("3.1", "string"),
    "sex": ("3.1", "string, M or F"),
    "birth_date": ("3.1", "Date"),
    "entry_date": ("3.1", "Date"),
    "maturity_date": ("3.1", "Date"),
    "prem_term_y": ("3.1", "Int (years)"),
    "sum_assured": ("3.1", "Money per policy"),
    "premium": ("3.1", "Money per policy, monthly, in advance"),
    "pols_if_init": ("3.1", "Count"),
}
# Section 3.2 derived fields: (symbol, formula on the inputs above, section, unit). Section 5's
# age last birthday compares (month, day) lexicographically; section 4.2's months_between is
# applied as section 4.3 says.
DERIVED: list[tuple[str, str, str, str]] = [
    (
        "age_at_entry",
        "=YEAR(entry_date)-YEAR(birth_date)"
        "-IF(OR(MONTH(entry_date)<MONTH(birth_date),"
        "AND(MONTH(entry_date)=MONTH(birth_date),DAY(entry_date)<DAY(birth_date))),1,0)",
        "5",
        "Int (years)",
    ),
    (
        "dur_m",
        "=12*(YEAR(valuation_date)-YEAR(entry_date))+(MONTH(valuation_date)-MONTH(entry_date))"
        "-IF(DAY(valuation_date)<DAY(entry_date),1,0)",
        "4.3",
        "Int (months)",
    ),
    (
        "term_m",
        "=12*(YEAR(maturity_date)-YEAR(entry_date))+(MONTH(maturity_date)-MONTH(entry_date))"
        "-IF(DAY(maturity_date)<DAY(entry_date),1,0)",
        "4.3",
        "Int (months)",
    ),
    ("rem_m", "=term_m-dur_m", "4.3", "Int (months)"),
    ("prem_term_m", "=12*prem_term_y", "4.3", "Int (months)"),
    ("prem_rem_m", "=MAX(prem_term_m-dur_m,0)", "4.3", "Int (months)"),
    ("dur_y_last", "=INT((term_m-1)/12)", "4.3", "Int (years)"),
]
MP_DERIVED_FIELDS = [name for name, _, _, _ in DERIVED]
DERIVED_META = {name: (section, unit) for name, _, section, unit in DERIVED}

INPUT_TITLE_ROW = 1
INPUT_FIRST_ROW = INPUT_TITLE_ROW + 1
DERIVED_TITLE_ROW = INPUT_FIRST_ROW + len(INPUT_FIELDS) + 1
DERIVED_FIRST_ROW = DERIVED_TITLE_ROW + 1
BASIS_TITLE_ROW = DERIVED_FIRST_ROW + len(DERIVED) + 1
BASIS_FIRST_ROW = BASIS_TITLE_ROW + 1
LAPSE_TITLE_ROW = BASIS_FIRST_ROW + len(BASIS) + 1
LAPSE_HEADER_ROW = LAPSE_TITLE_ROW + 1
COMM_TITLE_ROW = LAPSE_HEADER_ROW + len(LAPSE_TABLE) + 2
COMM_HEADER_ROW = COMM_TITLE_ROW + 1
MONTH_TITLE_ROW = COMM_HEADER_ROW + len(COMM_TABLE) + 2
MONTH_HEADER_ROW = MONTH_TITLE_ROW + 1
MONTH_FIRST_ROW = MONTH_HEADER_ROW + 1  # the row of t = 0; the row of t is MONTH_FIRST_ROW + t

# Policy sheet, monthly block: the section 11.1 columns in the section 7.1 evaluation order.
QUANTITIES: list[tuple[str, str, str]] = [
    ("t", "Int", "4.1"),
    ("dur_y", "Int", "5"),
    ("age", "Int", "5"),
    ("q_ann", "Prob annual", "6.1"),
    ("q_m", "Prob monthly", "6.1"),
    ("w_ann", "Prob annual", "6.2"),
    ("w_m", "Prob monthly", "6.2"),
    ("comm_rate", "Fraction", "6.6"),
    ("infl", "factor (state; infl(0) = 1)", "6.4"),
    ("deaths", "Count", "7.3"),
    ("lapses", "Count", "7.3"),
    ("maturities", "Count", "7.3"),
    ("pols_if", "Count (state)", "7.4"),
    ("premium_pp", "Money per policy", "8.1"),
    ("commission_pp", "Money per policy", "8.1"),
    ("maint_pp", "Money per policy", "8.1"),
    ("claim_exp_pp", "Money per policy", "8.1"),
    ("premium_income", "Money, start", "8.2"),
    ("commission", "Money, start", "8.2"),
    ("maint_expense", "Money, end", "8.2"),
    ("death_outgo", "Money, end", "8.2"),
    ("claim_expense", "Money, end", "8.2"),
    ("net_cf", "Money, reporting (no timing)", "8.3"),
    ("net_outgo", "Money, end", "8.3"),
    ("v", "factor", "9.1"),
    ("disc", "factor (state; disc(0) = 1)", "9.1"),
    ("pv_net_cf_t", "Money", "9.3"),
    ("pv_premiums_t", "Money", "9.3"),
    ("pv_commission_t", "Money", "9.3"),
    ("pv_maint_expense_t", "Money", "9.3"),
    ("pv_death_outgo_t", "Money", "9.3"),
    ("pv_claim_expense_t", "Money", "9.3"),
    ("reserve", "Money (state)", "10.2"),
    ("reserve_pp", "Money per policy (state)", "10.4"),
    ("reserve_floored", "Money (state)", "10.5"),
]

# Policy sheet, monthly block: the row-level check cells of section 12, after one blank column.
ROW_CHECKS: list[tuple[str, str]] = [
    ("H1", "pols_if >= 0 and non-increasing, exactly (t = 0 .. rem_m)"),
    (
        "H2",
        "rates in [0, 1] exactly; deaths + lapses + maturities <= pols_if(t-1) within tolerance",
    ),
    ("H4_lhs", "helper: roll-forward left-hand side"),
    ("H4", "roll-forward equals reserve(t) within tolerance"),
    ("H5b_term", "helper: month-s term of the H5b sum"),
    ("H5b_sum", "helper: sum of H5b_term over s = t+1 .. rem_m, ascending"),
    ("H5b", "reserve(t) * disc(t) equals the H5b sum within tolerance (t = 0 .. rem_m - 1)"),
    ("P1", "maturities exactly 0 except at rem_m, where it equals the survivors exactly"),
    ("P2", "premium, commission and lapse quantities exactly 0 after prem_rem_m; premium_income"),
    ("P4", "age in 18 .. 120; dur_y steps by 0 or 1"),
    ("P5_rhs", "helper: per-policy recursion right-hand side"),
    ("P5", "reserve_pp(t) equals the per-policy recursion within tolerance (scale sum_assured)"),
    ("P6", "reserve_floored >= 0 and >= reserve, exactly"),
    ("P10", "unit identities, exactly"),
]
FIRST_CHECK_COL = len(QUANTITIES) + 2  # one blank column after the quantities
LAST_COL = FIRST_CHECK_COL + len(ROW_CHECKS) - 1
COL = {name: get_column_letter(k + 1) for k, (name, _, _) in enumerate(QUANTITIES)}
COL.update({name: get_column_letter(FIRST_CHECK_COL + k) for k, (name, _) in enumerate(ROW_CHECKS)})

# Policy sheet, scalars block (section 11.2): the derived constants again, then Money and Count.
SCALAR_FIELDS = [
    *MP_DERIVED_FIELDS, "pv_premiums", "pv_commission", "pv_maint_expense", "pv_death_outgo",
    "pv_claim_expense", "pv_net_cf", "reserve_0", "reserve_pp_0", "reserve_floored_0",
    "maturities_n",
]  # fmt: skip
SCALAR_META: dict[str, tuple[str, str]] = {  # section, unit, for the Money and Count scalars
    "pv_premiums": ("9.3", "Money"),
    "pv_commission": ("9.3", "Money"),
    "pv_maint_expense": ("9.3", "Money"),
    "pv_death_outgo": ("9.3", "Money"),
    "pv_claim_expense": ("9.3", "Money"),
    "pv_net_cf": ("9.3", "Money"),
    "reserve_0": ("10.2", "Money: reserve(0)"),
    "reserve_pp_0": ("10.4", "Money per policy: reserve_pp(0)"),
    "reserve_floored_0": ("10.5", "Money: reserve_floored(0)"),
    "maturities_n": ("7.3", "Count: maturities(rem_m), the survivors to the end of the term"),
}
PV_SUMS = {
    "pv_premiums": "pv_premiums_t",
    "pv_commission": "pv_commission_t",
    "pv_maint_expense": "pv_maint_expense_t",
    "pv_death_outgo": "pv_death_outgo_t",
    "pv_claim_expense": "pv_claim_expense_t",
}
SCALAR_TOTALS = [
    "pv_premiums", "pv_commission", "pv_maint_expense", "pv_death_outgo", "pv_claim_expense",
    "pv_net_cf", "reserve_0", "reserve_floored_0",
]  # fmt: skip

# Policy sheet, check block (section 12): the section 1 scales and two helper sums, each in one
# named cell, then one ``<id>_ok`` per (sheet) invariant, ``rows_ok`` and ``ok``.
CHECK_HELPERS: dict[str, str] = {
    "scale_money": "Money (in-force) scale of section 1: sum_assured * pols_if_init",
    "scale_pp": "Money per policy scale of section 1: sum_assured",
    "scale_count": "Count scale of section 1: pols_if_init",
    "exits_sum": "helper (P8): sum over t = 1 .. rem_m of deaths + lapses + maturities",
    "pv_net_cf_t_sum": "helper (P9): sum over t = 1 .. rem_m of pv_net_cf_t",
}
SHEET_INVARIANTS = [
    "H1", "H2", "H3", "H4", "H5a", "H5b", "P1", "P2", "P3", "P4", "P5", "P6", "P8", "P9", "P10",
]  # fmt: skip
OK_CELLS = [f"{inv}_ok" for inv in SHEET_INVARIANTS] + ["rows_ok", "ok"]
OK_MEANING: dict[str, str] = {
    "H1_ok": "every H1 row cell (t = 0 .. rem_m)",
    "H2_ok": "every H2 row cell (t = 1 .. rem_m)",
    "H3_ok": "reserve(rem_m) = 0 and pols_if(rem_m) = 0, exactly",
    "H4_ok": "every H4 row cell (t = 1 .. rem_m)",
    "H5a_ok": "reserve(0) equals -pv_net_cf within tolerance (scale_money)",
    "H5b_ok": "every H5b row cell (t = 0 .. rem_m - 1)",
    "P1_ok": "every P1 row cell (t = 1 .. rem_m)",
    "P2_ok": "every P2 row cell (t = 1 .. rem_m)",
    "P3_ok": "the number of months with premium_pp <> 0 equals prem_rem_m",
    "P4_ok": "every P4 row cell (t = 1 .. rem_m)",
    "P5_ok": "every P5 row cell (t = 0 .. rem_m - 1)",
    "P6_ok": "every P6 row cell (t = 0 .. rem_m)",
    "P8_ok": "exits_sum equals pols_if_init within tolerance (scale_count)",
    "P9_ok": "pv_net_cf equals pv_net_cf_t_sum within tolerance (scale_money)",
    "P10_ok": "every P10 row cell (t = 1 .. rem_m)",
    "rows_ok": "layout: the last t of the monthly block equals the formula-derived rem_m",
    "ok": "every cell above",
}

# summary sheet: the three blocks, top to bottom.
SUMMARY_SCALAR_COLUMNS = ["policy_id", *SCALAR_FIELDS]
SCOL = {name: get_column_letter(k + 1) for k, name in enumerate(SUMMARY_SCALAR_COLUMNS)}
SUMMARY_CHECK_COLUMNS = ["policy_id", *OK_CELLS]
CCOL = {name: get_column_letter(k + 1) for k, name in enumerate(SUMMARY_CHECK_COLUMNS)}
# The Money and Count monthly columns of 11.1 (per-policy Money is not summed).
STATE_TOTALS = ["pols_if", "reserve", "reserve_floored"]
TOTAL_COLUMNS = [
    "deaths", "lapses", "maturities", "pols_if", "premium_income", "commission",
    "maint_expense", "death_outgo", "claim_expense", "net_cf", "net_outgo", "pv_net_cf_t",
    "pv_premiums_t", "pv_commission_t", "pv_maint_expense_t", "pv_death_outgo_t",
    "pv_claim_expense_t", "reserve", "reserve_floored",
]  # fmt: skip
TCOL = {name: get_column_letter(k + 2) for k, name in enumerate(TOTAL_COLUMNS)}

SUMMARY_SCALAR_TITLE_ROW = 1
SUMMARY_SCALAR_HEADER_ROW = SUMMARY_SCALAR_TITLE_ROW + 1
SUMMARY_SCALAR_FIRST_ROW = SUMMARY_SCALAR_HEADER_ROW + 1
SUMMARY_TOTAL_ROW = SUMMARY_SCALAR_FIRST_ROW + len(MODEL_POINTS)
SUMMARY_CHECK_TITLE_ROW = SUMMARY_TOTAL_ROW + 2
SUMMARY_CHECK_HEADER_ROW = SUMMARY_CHECK_TITLE_ROW + 1
SUMMARY_CHECK_FIRST_ROW = SUMMARY_CHECK_HEADER_ROW + 1
SUMMARY_ALL_OK_ROW = SUMMARY_CHECK_FIRST_ROW + len(MODEL_POINTS) + 1
SUMMARY_TOTALS_T_TITLE_ROW = SUMMARY_ALL_OK_ROW + 2
SUMMARY_TOTALS_T_HEADER_ROW = SUMMARY_TOTALS_T_TITLE_ROW + 1
SUMMARY_TOTALS_T_FIRST_ROW = SUMMARY_TOTALS_T_HEADER_ROW + 1

# mortality sheet: the 6.1 parameters, then the table.
MORT_PARAM_FIRST_ROW = 2
MORT_TITLE_ROW = MORT_PARAM_FIRST_ROW + len(MORT_PARAMS) + 1
MORT_HEADER_ROW = MORT_TITLE_ROW + 1
MORT_FIRST_ROW = MORT_HEADER_ROW + 1

BOLD = Font(bold=True)


@dataclass(frozen=True)
class Blocks:
    """Rows of the blocks of one policy sheet that follow its monthly block, from its ``rem_m``."""

    rem_m: int

    @property
    def last(self) -> int:
        """The row of ``t = rem_m``, the last row of the monthly block."""
        return MONTH_FIRST_ROW + self.rem_m

    @property
    def scalar_title(self) -> int:
        return self.last + 2

    def scalar_row(self, name: str) -> int:
        return self.scalar_title + 1 + SCALAR_FIELDS.index(name)

    @property
    def check_title(self) -> int:
        return self.scalar_title + len(SCALAR_FIELDS) + 2

    def helper_row(self, name: str) -> int:
        return self.check_title + 1 + list(CHECK_HELPERS).index(name)

    def ok_row(self, name: str) -> int:
        return self.check_title + len(CHECK_HELPERS) + 2 + OK_CELLS.index(name)


class Row:
    """Cell addresses of row ``t`` of the monthly block by spec symbol: ``Row(1).q_m == "E53"``."""

    def __init__(self, t: int) -> None:
        self.r = MONTH_FIRST_ROW + t

    def __getattr__(self, name: str) -> str:
        return f"{COL[name]}{self.r}"


def cell_at(name: str, t: int) -> str:
    """Absolute same-sheet address of monthly column ``name`` in row ``t``."""
    return f"${COL[name]}${MONTH_FIRST_ROW + t}"


def column_range(name: str, first: int, last: int) -> str:
    """Absolute same-sheet range of monthly column ``name`` over rows ``t = first .. last``."""
    return f"{cell_at(name, first)}:{cell_at(name, last)}"


def define(wb: openpyxl.Workbook, name: str, ref: str) -> None:
    """Add a workbook-level defined name for ``ref``."""
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def define_local(ws: Worksheet, name: str, ref: str) -> None:
    """Add a defined name scoped to ``ws``: formulas on that sheet resolve it, others do not."""
    ws.defined_names[name] = DefinedName(name, attr_text=ref)


def title(ws: Worksheet, row: int, text: str) -> None:
    ws[f"A{row}"] = text
    ws[f"A{row}"].font = BOLD


def set_widths(ws: Worksheet, width: float, upto: int) -> None:
    for k in range(1, upto + 1):
        ws.column_dimensions[get_column_letter(k)].width = width


# ----------------------------------------------------------------------------------------------
# readme
# ----------------------------------------------------------------------------------------------


def write_readme(wb: openpyxl.Workbook) -> None:
    ws = wb.active
    assert ws is not None
    ws.title = "readme"
    b0 = Blocks(0)  # the rows after the monthly block, for a policy with rem_m = 0
    lines = [
        "Halley term assurance: reference workbook for docs/spec/term-assurance.md (version 0.1).",
        "Built from the specification alone: typed cells are its reference inputs (sections 3.3,"
        " 6, 13); every other cell is a formula transcribing one formula of the spec.",
        "",
        "Sheets",
        "mortality: the 6.1 parameters A and B by sex and c, and the mortality table q_ann(sex, x)"
        " materialised for ages 18 to 120 from the 6.1 formula in its stated order (section 11,"
        " materialised-block option). The one block shared by the policy sheets: each q_ann cell"
        " looks it up by INDEX/MATCH on (sex, age).",
        "P01 .. P10: one sheet per policy, named by policy_id, with the section 11 blocks in order"
        " from the top: inputs (3.1, plus the run-level valuation_date of 3.3), typed; derived"
        " constants (3.2, 4.3, 5) by formula; basis constants (6) with i_m, v_m, e_m by formula;"
        " lapse table (6.2); commission table (6.6); the monthly block (7.1, 11.1) with a header"
        " row and rows t = 0 .. rem_m, the row checks of section 12 to its right; scalars (11.2);"
        " check cells (12). Every formula reads this sheet's own cells, through sheet-scoped"
        " defined names (sex, premium, rem_m, i_m, lapse_table, scale_money, ...) or same-sheet"
        " absolute references, except the q_ann lookup on mortality.",
        "summary: the 11.2 scalars per policy by reference, the 11.3 scalar totals in the row"
        " labelled total (plain SUMs in ascending policy_id order), the <id>_ok cells per policy"
        " by reference with ok per policy and all_ok (defined name all_ok), and the 11.3 totals"
        " by t = 0 .. 480 over the ten policies, policies beyond their own rem_m read as zero.",
        "",
        "Conventions",
        "q_ann is looked up from the materialised (sex, age, q_ann) block on mortality, ages 18 to"
        " 120, each cell the section 6.1 formula in its stated order (section 11 option 2); the"
        " 6.1 parameters are typed on mortality (the table reads them) and on every policy sheet"
        " (the basis block of section 11).",
        "Row t = 0 holds the state variables only (pols_if, infl, disc, reserve, reserve_pp,"
        " reserve_floored); month quantities are blank there (section 4.1).",
        "Tolerance (section 1): |a - b| <= 1e-9 * max(|a|, |b|, scale); inequalities use"
        " a <= b + 1e-9 * scale; 'exactly' is a plain = comparison. The scales are the check"
        " block's scale_money, scale_pp and scale_count, from the sheet's own inputs.",
        "P7 and P11 are engine-only (rows beyond rem_m, batch independence) and have no cell.",
        "The spec's c is the defined name gm_c because Excel reserves the names R and C.",
        "rows_ok on a policy sheet compares the last t of the monthly block (laid out by the"
        " builder) with the formula-derived rem_m; it is a layout check, not a spec invariant.",
        "",
        "Policy sheet blocks (block, rows; column A label or t, B value, C section, D unit)",
    ]
    blocks = [
        (
            "inputs (3.1, 3.3), typed",
            f"{INPUT_FIRST_ROW} to {INPUT_FIRST_ROW + len(INPUT_FIELDS) - 1}",
        ),
        (
            "derived constants (3.2, 4.3, 5), by formula",
            f"{DERIVED_FIRST_ROW} to {DERIVED_FIRST_ROW + len(DERIVED) - 1}",
        ),
        (
            "basis constants (6), i_m, v_m, e_m by formula",
            f"{BASIS_FIRST_ROW} to {BASIS_FIRST_ROW + len(BASIS) - 1}",
        ),
        ("lapse table (6.2)", f"{LAPSE_HEADER_ROW + 1} to {LAPSE_HEADER_ROW + len(LAPSE_TABLE)}"),
        ("commission table (6.6)", f"{COMM_HEADER_ROW + 1} to {COMM_HEADER_ROW + len(COMM_TABLE)}"),
        (
            "monthly block (7.1, 11.1) and row checks (12)",
            f"header {MONTH_HEADER_ROW}; row of t is {MONTH_FIRST_ROW} + t",
        ),
        ("scalars (11.2)", f"from {b0.scalar_row(SCALAR_FIELDS[0])} + rem_m"),
        (
            "check cells (12): scales and helpers, then <id>_ok, rows_ok, ok",
            f"from {b0.helper_row('scale_money')} + rem_m; <id>_ok from"
            f" {b0.ok_row(OK_CELLS[0])} + rem_m",
        ),
    ]
    for k, line in enumerate(lines, start=1):
        ws.cell(row=k, column=1, value=line)
    row = len(lines) + 1
    for name, rows in blocks:
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=rows)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Policy sheet monthly columns (symbol, unit, spec section)")
    row += 1
    for name, unit, section in QUANTITIES:
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=unit)
        ws.cell(row=row, column=3, value=section)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Policy sheet row check columns (section 12)")
    row += 1
    for name, meaning in ROW_CHECKS:
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=meaning)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Policy sheet check cells (section 12), repeated on summary")
    row += 1
    for name in OK_CELLS:
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=OK_MEANING[name])
        row += 1
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 60


# ----------------------------------------------------------------------------------------------
# mortality: the shared materialised table of section 6.1
# ----------------------------------------------------------------------------------------------


def write_mortality(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("mortality")
    title(
        ws, 1, "Mortality parameters (section 6.1), each in one cell: the table below reads these"
    )
    for k, (symbol, value, unit, section) in enumerate(MORT_PARAMS):
        row = MORT_PARAM_FIRST_ROW + k
        ws[f"A{row}"] = symbol
        ws[f"B{row}"] = value
        ws[f"C{row}"] = section
        ws[f"D{row}"] = unit
        define_local(ws, EXCEL_NAME.get(symbol, symbol), f"{ws.title}!$B${row}")

    title(
        ws,
        MORT_TITLE_ROW,
        "Mortality table q_ann(sex, x) (section 6.1): 1 - exp(-(A + B * c^x * (c - 1) / ln(c))),"
        " ages 18 to 120; shared by every policy sheet, which looks it up by (sex, age)",
    )
    ws[f"A{MORT_HEADER_ROW}"] = "x"
    ws[f"B{MORT_HEADER_ROW}"] = "M"
    ws[f"C{MORT_HEADER_ROW}"] = "F"
    for k, age in enumerate(MORT_AGES):
        row = MORT_FIRST_ROW + k
        ws[f"A{row}"] = age
        ws[f"B{row}"] = f"=1-EXP(-(A_M+B_M*gm_c^$A{row}*(gm_c-1)/LN(gm_c)))"
        ws[f"C{row}"] = f"=1-EXP(-(A_F+B_F*gm_c^$A{row}*(gm_c-1)/LN(gm_c)))"
    last = MORT_FIRST_ROW + len(MORT_AGES) - 1
    define(wb, "mort_age", f"{ws.title}!$A${MORT_FIRST_ROW}:$A${last}")
    define(wb, "mort_sex", f"{ws.title}!$B${MORT_HEADER_ROW}:$C${MORT_HEADER_ROW}")
    define(wb, "mort_table", f"{ws.title}!$B${MORT_FIRST_ROW}:$C${last}")
    set_widths(ws, 16, 4)


# ----------------------------------------------------------------------------------------------
# One sheet per policy
# ----------------------------------------------------------------------------------------------


def write_policy(wb: openpyxl.Workbook, mp: ModelPoint) -> None:
    ws = wb.create_sheet(mp.policy_id)
    blocks = Blocks(mp.rem_m)
    write_inputs_block(ws, mp)
    write_derived_block(ws)
    write_basis_block(ws)
    write_monthly_block(ws, mp.rem_m)
    write_scalars_block(ws, blocks)
    write_checks_block(ws, blocks)
    ws.freeze_panes = "B1"
    set_widths(ws, 15, LAST_COL)
    ws.column_dimensions["A"].width = 20


def write_inputs_block(ws: Worksheet, mp: ModelPoint) -> None:
    """Section 3.1 inputs of this model point (section 13) and the 3.3 valuation date, typed."""
    title(
        ws,
        INPUT_TITLE_ROW,
        "Inputs (section 3.1, from section 13) and the run-level valuation_date (3.3), typed",
    )
    for k, name in enumerate(INPUT_FIELDS):
        row = INPUT_FIRST_ROW + k
        value = VALUATION_DATE if name == "valuation_date" else getattr(mp, name)
        section, unit = INPUT_META[name]
        ws[f"A{row}"] = name
        ws[f"B{row}"] = value
        if isinstance(value, dt.date):
            ws[f"B{row}"].number_format = "yyyy-mm-dd"
        ws[f"C{row}"] = section
        ws[f"D{row}"] = unit
        if name != "policy_id":
            define_local(ws, name, f"{ws.title}!$B${row}")


def write_derived_block(ws: Worksheet) -> None:
    """Section 3.2 derived fields by the formulas of sections 4.3 and 5, on the inputs above."""
    title(ws, DERIVED_TITLE_ROW, "Derived constants (sections 3.2, 4.3 and 5), by formula")
    for k, (name, formula, section, unit) in enumerate(DERIVED):
        row = DERIVED_FIRST_ROW + k
        ws[f"A{row}"] = name
        ws[f"B{row}"] = formula
        ws[f"C{row}"] = section
        ws[f"D{row}"] = unit
        define_local(ws, name, f"{ws.title}!$B${row}")


def write_basis_block(ws: Worksheet) -> None:
    """Section 6: the constants, each in one cell, then the lapse and commission tables."""
    title(ws, BASIS_TITLE_ROW, "Basis constants (section 6), each in one cell")
    for k, (symbol, value, unit, section) in enumerate(BASIS):
        row = BASIS_FIRST_ROW + k
        ws[f"A{row}"] = symbol
        ws[f"B{row}"] = value
        ws[f"C{row}"] = section
        ws[f"D{row}"] = unit
        define_local(ws, EXCEL_NAME.get(symbol, symbol), f"{ws.title}!$B${row}")

    title(ws, LAPSE_TITLE_ROW, "Lapse table (section 6.2): key k = min(dur_y, 5)")
    ws[f"A{LAPSE_HEADER_ROW}"] = "k"
    ws[f"B{LAPSE_HEADER_ROW}"] = "lapse_table"
    ws[f"D{LAPSE_HEADER_ROW}"] = "Prob, annual"
    for k, w_ann in enumerate(LAPSE_TABLE):
        ws[f"A{LAPSE_HEADER_ROW + 1 + k}"] = k
        ws[f"B{LAPSE_HEADER_ROW + 1 + k}"] = w_ann
    first, last = LAPSE_HEADER_ROW + 1, LAPSE_HEADER_ROW + len(LAPSE_TABLE)
    define_local(ws, "lapse_key", f"{ws.title}!$A${first}:$A${last}")
    define_local(ws, "lapse_table", f"{ws.title}!$B${first}:$B${last}")

    title(ws, COMM_TITLE_ROW, "Commission table (section 6.6): key k = min(dur_y, 1)")
    ws[f"A{COMM_HEADER_ROW}"] = "k"
    ws[f"B{COMM_HEADER_ROW}"] = "comm_table"
    ws[f"D{COMM_HEADER_ROW}"] = "Fraction"
    for k, rate in enumerate(COMM_TABLE):
        ws[f"A{COMM_HEADER_ROW + 1 + k}"] = k
        ws[f"B{COMM_HEADER_ROW + 1 + k}"] = rate
    first, last = COMM_HEADER_ROW + 1, COMM_HEADER_ROW + len(COMM_TABLE)
    define_local(ws, "comm_key", f"{ws.title}!$A${first}:$A${last}")
    define_local(ws, "comm_table", f"{ws.title}!$B${first}:$B${last}")


def write_monthly_block(ws: Worksheet, rem_m: int) -> None:
    """Sections 7.1 and 11.1: the header row, rows ``t = 0 .. rem_m`` and the row checks."""
    title(
        ws,
        MONTH_TITLE_ROW,
        "Monthly block (sections 7.1 and 11.1): rows t = 0 .. rem_m, one column per quantity in"
        " the 7.1 order; the row-level check cells of section 12 follow the blank column",
    )
    for name, _, _ in QUANTITIES:
        ws[f"{COL[name]}{MONTH_HEADER_ROW}"] = name
        ws[f"{COL[name]}{MONTH_HEADER_ROW}"].font = BOLD
    for name, _ in ROW_CHECKS:
        ws[f"{COL[name]}{MONTH_HEADER_ROW}"] = name
        ws[f"{COL[name]}{MONTH_HEADER_ROW}"].font = BOLD
    for t in range(rem_m + 1):
        write_policy_row(ws, t, rem_m)
    for t in range(rem_m + 1):
        write_check_row(ws, t, rem_m)


def write_policy_row(ws: Worksheet, t: int, rem_m: int) -> None:
    """Row ``t`` of the monthly block (section 7.1 steps 2 to 4)."""
    cur, prv, nxt = Row(t), Row(t - 1), Row(t + 1)
    ws[cur.t] = t
    if t == 0:
        # Step 2: state variables take their initial values; month quantities stay blank.
        ws[cur.infl] = 1
        ws[cur.pols_if] = "=pols_if_init"
        ws[cur.disc] = 1
    else:
        # Step 3: each row from its own t and row t-1 only, columns left to right.
        ws[cur.dur_y] = f"=MIN(INT((dur_m+{cur.t}-1)/12),dur_y_last)"
        ws[cur.age] = f"=age_at_entry+{cur.dur_y}"
        ws[cur.q_ann] = f"=INDEX(mort_table,MATCH({cur.age},mort_age,0),MATCH(sex,mort_sex,0))"
        ws[cur.q_m] = f"=IF({cur.t}<=rem_m,1-(1-{cur.q_ann})^(1/12),0)"
        ws[cur.w_ann] = (
            f"=IF({cur.t}<=prem_rem_m,INDEX(lapse_table,MATCH(MIN({cur.dur_y},5),lapse_key,0)),0)"
        )
        ws[cur.w_m] = f"=1-(1-{cur.w_ann})^(1/12)"
        ws[cur.comm_rate] = f"=INDEX(comm_table,MATCH(MIN({cur.dur_y},1),comm_key,0))"
        ws[cur.infl] = f"=(1+e_m)^{cur.t}"
        ws[cur.deaths] = f"={prv.pols_if}*{cur.q_m}"
        ws[cur.lapses] = f"=({prv.pols_if}-{cur.deaths})*{cur.w_m}"
        ws[cur.maturities] = f"=IF({cur.t}=rem_m,{prv.pols_if}-{cur.deaths}-{cur.lapses},0)"
        ws[cur.pols_if] = f"={prv.pols_if}-{cur.deaths}-{cur.lapses}-{cur.maturities}"
        ws[cur.premium_pp] = f"=IF({cur.t}<=prem_rem_m,premium,0)"
        ws[cur.commission_pp] = f"={cur.premium_pp}*{cur.comm_rate}"
        ws[cur.maint_pp] = f"=maint_base*{cur.infl}"
        ws[cur.claim_exp_pp] = f"=claim_exp_base*{cur.infl}"
        ws[cur.premium_income] = f"={prv.pols_if}*{cur.premium_pp}"
        ws[cur.commission] = f"={cur.premium_income}*{cur.comm_rate}"
        ws[cur.maint_expense] = f"={prv.pols_if}*{cur.maint_pp}"
        ws[cur.death_outgo] = f"={cur.deaths}*sum_assured"
        ws[cur.claim_expense] = f"={cur.deaths}*{cur.claim_exp_pp}"
        ws[cur.net_cf] = (
            f"={cur.premium_income}-{cur.commission}-{cur.maint_expense}"
            f"-{cur.death_outgo}-{cur.claim_expense}"
        )
        ws[cur.net_outgo] = (
            f"={cur.death_outgo}+{cur.claim_expense}+{cur.maint_expense}"
            f"-({cur.premium_income}-{cur.commission})*(1+i_m)"
        )
        ws[cur.v] = "=v_m"
        ws[cur.disc] = f"=v_m^{cur.t}"
        ws[cur.pv_net_cf_t] = (
            f"=({cur.premium_income}-{cur.commission})*{prv.disc}"
            f"-({cur.maint_expense}+{cur.death_outgo}+{cur.claim_expense})*{cur.disc}"
        )
        ws[cur.pv_premiums_t] = f"={cur.premium_income}*{prv.disc}"
        ws[cur.pv_commission_t] = f"={cur.commission}*{prv.disc}"
        ws[cur.pv_maint_expense_t] = f"={cur.maint_expense}*{cur.disc}"
        ws[cur.pv_death_outgo_t] = f"={cur.death_outgo}*{cur.disc}"
        ws[cur.pv_claim_expense_t] = f"={cur.claim_expense}*{cur.disc}"
    # Step 4: the reserve by backward recursion from reserve(rem_m) = 0, then the two derived.
    if t == rem_m:
        ws[cur.reserve] = 0
    else:
        ws[cur.reserve] = f"=({nxt.reserve}+{nxt.net_outgo})*{nxt.v}"
    ws[cur.reserve_pp] = f"=IF({cur.pols_if}>0,{cur.reserve}/{cur.pols_if},0)"
    ws[cur.reserve_floored] = f"=IF({cur.reserve}>0,{cur.reserve},0)"


def write_check_row(ws: Worksheet, t: int, rem_m: int) -> None:
    """Row ``t`` of the section 12 row-level check cells."""
    cur, prv, nxt = Row(t), Row(t - 1), Row(t + 1)
    tol = "1E-9*"
    if t == 0:
        ws[cur.H1] = f"={cur.pols_if}>=0"
    else:
        ws[cur.H1] = f"=AND({cur.pols_if}>=0,{cur.pols_if}<={prv.pols_if})"
        ws[cur.H2] = (
            f"=AND({cur.q_ann}>=0,{cur.q_ann}<=1,{cur.q_m}>=0,{cur.q_m}<=1,"
            f"{cur.w_ann}>=0,{cur.w_ann}<=1,{cur.w_m}>=0,{cur.w_m}<=1,"
            f"{cur.deaths}+{cur.lapses}+{cur.maturities}<={prv.pols_if}+{tol}scale_count)"
        )
        ws[cur.H4_lhs] = (
            f"=({prv.reserve}+{cur.premium_income}-{cur.commission})*(1+i_m)"
            f"-{cur.maint_expense}-{cur.death_outgo}-{cur.claim_expense}"
        )
        ws[cur.H4] = (
            f"=ABS({cur.H4_lhs}-{cur.reserve})"
            f"<={tol}MAX(ABS({cur.H4_lhs}),ABS({cur.reserve}),scale_money)"
        )
        ws[cur.H5b_term] = (
            f"=({cur.death_outgo}+{cur.claim_expense}+{cur.maint_expense})*{cur.disc}"
            f"-({cur.premium_income}-{cur.commission})*{prv.disc}"
        )
        ws[cur.P1] = (
            f"=IF({cur.t}<>rem_m,{cur.maturities}=0,"
            f"{cur.maturities}={prv.pols_if}-{cur.deaths}-{cur.lapses})"
        )
        ws[cur.P2] = (
            f"=IF({cur.t}>prem_rem_m,"
            f"AND({cur.premium_income}=0,{cur.commission}=0,{cur.w_ann}=0,{cur.w_m}=0,"
            f"{cur.lapses}=0),{cur.premium_income}={prv.pols_if}*premium)"
        )
        age_ok = f"{cur.age}>=18,{cur.age}<=120"
        if t == 1:
            ws[cur.P4] = f"=AND({age_ok})"
        else:
            step = f"{cur.dur_y}-{prv.dur_y}"
            ws[cur.P4] = f"=AND({age_ok},OR({step}=0,{step}=1))"
        ws[cur.P10] = (
            f"=AND({cur.death_outgo}={cur.deaths}*sum_assured,"
            f"{cur.premium_income}={prv.pols_if}*{cur.premium_pp},"
            f"{cur.maint_expense}={prv.pols_if}*{cur.maint_pp})"
        )
    if t < rem_m:
        ws[cur.H5b_sum] = f"=SUM({nxt.H5b_term}:{cell_at('H5b_term', rem_m)})"
        lhs = f"{cur.reserve}*{cur.disc}"
        ws[cur.H5b] = (
            f"=ABS({lhs}-{cur.H5b_sum})<={tol}MAX(ABS({lhs}),ABS({cur.H5b_sum}),scale_money)"
        )
        ws[cur.P5_rhs] = (
            f"=-({nxt.premium_pp}-{nxt.commission_pp})"
            f"+v_m*({nxt.q_m}*(sum_assured+{nxt.claim_exp_pp})+{nxt.maint_pp}"
            f"+(1-{nxt.q_m})*(1-{nxt.w_m})*{nxt.reserve_pp})"
        )
        ws[cur.P5] = (
            f"=ABS({cur.reserve_pp}-{cur.P5_rhs})"
            f"<={tol}MAX(ABS({cur.reserve_pp}),ABS({cur.P5_rhs}),scale_pp)"
        )
    ws[cur.P6] = f"=AND({cur.reserve_floored}>=0,{cur.reserve_floored}>={cur.reserve})"


def write_scalars_block(ws: Worksheet, b: Blocks) -> None:
    """Section 11.2 (7.1 step 5): the derived constants again, the PV sums, the time-0 state."""
    title(
        ws,
        b.scalar_title,
        "Scalars (section 11.2): the derived constants, the PV sums over t = 1 .. rem_m in"
        " ascending t, the time-0 state variables and maturities(rem_m)",
    )
    for name in SCALAR_FIELDS:
        row = b.scalar_row(name)
        ws[f"A{row}"] = name
        if name in DERIVED_META:
            ws[f"B{row}"] = f"={name}"
            section, unit = DERIVED_META[name]
        else:
            section, unit = SCALAR_META[name]
            define_local(ws, name, f"{ws.title}!$B${row}")
        ws[f"C{row}"] = section
        ws[f"D{row}"] = unit
    # Section 9.3: each PV is the sum of its column over t = 1 .. rem_m, then pv_net_cf from the
    # five, left to right.
    for name, column in PV_SUMS.items():
        ws[f"B{b.scalar_row(name)}"] = f"=SUM({column_range(column, 1, b.rem_m)})"
    ws[f"B{b.scalar_row('pv_net_cf')}"] = "=" + "-".join(f"B{b.scalar_row(n)}" for n in PV_SUMS)
    ws[f"B{b.scalar_row('reserve_0')}"] = f"={cell_at('reserve', 0)}"
    ws[f"B{b.scalar_row('reserve_pp_0')}"] = f"={cell_at('reserve_pp', 0)}"
    ws[f"B{b.scalar_row('reserve_floored_0')}"] = f"={cell_at('reserve_floored', 0)}"
    ws[f"B{b.scalar_row('maturities_n')}"] = f"={cell_at('maturities', b.rem_m)}"


def write_checks_block(ws: Worksheet, b: Blocks) -> None:
    """Section 12 (7.1 step 5): scales and helper sums, then ``<id>_ok``, ``rows_ok``, ``ok``."""
    rem_m, tol = b.rem_m, "1E-9*"
    title(
        ws,
        b.check_title,
        "Check cells (section 12): the section 1 scales and two helper sums, then one <id>_ok per"
        " (sheet) invariant, TRUE when every row cell and scalar cell of that invariant is TRUE",
    )
    helpers = {
        "scale_money": "=sum_assured*pols_if_init",
        "scale_pp": "=sum_assured",
        "scale_count": "=pols_if_init",
        "exits_sum": (
            f"=SUMPRODUCT({column_range('deaths', 1, rem_m)}+{column_range('lapses', 1, rem_m)}"
            f"+{column_range('maturities', 1, rem_m)})"
        ),
        "pv_net_cf_t_sum": f"=SUM({column_range('pv_net_cf_t', 1, rem_m)})",
    }
    for name, meaning in CHECK_HELPERS.items():
        row = b.helper_row(name)
        ws[f"A{row}"] = name
        ws[f"B{row}"] = helpers[name]
        ws[f"D{row}"] = meaning
        define_local(ws, name, f"{ws.title}!$B${row}")
    ok = {
        "H1_ok": f"=AND({column_range('H1', 0, rem_m)})",
        "H2_ok": f"=AND({column_range('H2', 1, rem_m)})",
        "H3_ok": f"=AND({cell_at('reserve', rem_m)}=0,{cell_at('pols_if', rem_m)}=0)",
        "H4_ok": f"=AND({column_range('H4', 1, rem_m)})",
        "H5a_ok": (
            f"=ABS(reserve_0-(-pv_net_cf))<={tol}MAX(ABS(reserve_0),ABS(pv_net_cf),scale_money)"
        ),
        "H5b_ok": f"=AND({column_range('H5b', 0, rem_m - 1)})",
        "P1_ok": f"=AND({column_range('P1', 1, rem_m)})",
        "P2_ok": f"=AND({column_range('P2', 1, rem_m)})",
        "P3_ok": f"=SUMPRODUCT(--({column_range('premium_pp', 1, rem_m)}<>0))=prem_rem_m",
        "P4_ok": f"=AND({column_range('P4', 1, rem_m)})",
        "P5_ok": f"=AND({column_range('P5', 0, rem_m - 1)})",
        "P6_ok": f"=AND({column_range('P6', 0, rem_m)})",
        "P8_ok": (
            f"=ABS(exits_sum-scale_count)<={tol}MAX(ABS(exits_sum),ABS(scale_count),scale_count)"
        ),
        "P9_ok": (
            f"=ABS(pv_net_cf-pv_net_cf_t_sum)"
            f"<={tol}MAX(ABS(pv_net_cf),ABS(pv_net_cf_t_sum),scale_money)"
        ),
        "P10_ok": f"=AND({column_range('P10', 1, rem_m)})",
        "rows_ok": f"=MAX({column_range('t', 0, rem_m)})=rem_m",
    }
    ok["ok"] = f"=AND($B${b.ok_row(OK_CELLS[0])}:$B${b.ok_row('rows_ok')})"
    for name in OK_CELLS:
        row = b.ok_row(name)
        ws[f"A{row}"] = name
        ws[f"B{row}"] = ok[name]
        ws[f"D{row}"] = OK_MEANING[name]
    ws[f"A{b.ok_row('ok')}"].font = BOLD


# ----------------------------------------------------------------------------------------------
# summary
# ----------------------------------------------------------------------------------------------


def write_summary(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("summary")
    blocks = {mp.policy_id: Blocks(mp.rem_m) for mp in MODEL_POINTS}
    policy_id_row = INPUT_FIRST_ROW + INPUT_FIELDS.index("policy_id")

    # Section 11.2 scalars per policy, by reference; section 11.3 scalar totals as plain SUMs.
    title(
        ws,
        SUMMARY_SCALAR_TITLE_ROW,
        "Scalars (section 11.2) per policy, by reference to the policy sheets; the row labelled"
        " total holds the section 11.3 scalar totals, plain SUMs in ascending policy_id order",
    )
    for name in SUMMARY_SCALAR_COLUMNS:
        ws[f"{SCOL[name]}{SUMMARY_SCALAR_HEADER_ROW}"] = name
        ws[f"{SCOL[name]}{SUMMARY_SCALAR_HEADER_ROW}"].font = BOLD
    for k, mp in enumerate(MODEL_POINTS):
        row, b, pid = SUMMARY_SCALAR_FIRST_ROW + k, blocks[mp.policy_id], mp.policy_id
        ws[f"{SCOL['policy_id']}{row}"] = f"={pid}!$B${policy_id_row}"
        for name in SCALAR_FIELDS:
            ws[f"{SCOL[name]}{row}"] = f"={pid}!$B${b.scalar_row(name)}"
    ws[f"A{SUMMARY_TOTAL_ROW}"] = "total"
    ws[f"A{SUMMARY_TOTAL_ROW}"].font = BOLD
    for name in SCALAR_TOTALS:
        col = SCOL[name]
        ws[f"{col}{SUMMARY_TOTAL_ROW}"] = (
            f"=SUM({col}{SUMMARY_SCALAR_FIRST_ROW}:{col}{SUMMARY_TOTAL_ROW - 1})"
        )

    # Section 12 check cells per policy, by reference, and all_ok.
    title(
        ws,
        SUMMARY_CHECK_TITLE_ROW,
        "Check cells (section 12) per policy, by reference to the policy sheets: <id>_ok per"
        " (sheet) invariant, rows_ok, ok per policy, and all_ok over the ten policies",
    )
    for name in SUMMARY_CHECK_COLUMNS:
        ws[f"{CCOL[name]}{SUMMARY_CHECK_HEADER_ROW}"] = name
        ws[f"{CCOL[name]}{SUMMARY_CHECK_HEADER_ROW}"].font = BOLD
    for k, mp in enumerate(MODEL_POINTS):
        row, b, pid = SUMMARY_CHECK_FIRST_ROW + k, blocks[mp.policy_id], mp.policy_id
        ws[f"{CCOL['policy_id']}{row}"] = f"={pid}!$B${policy_id_row}"
        for name in OK_CELLS:
            ws[f"{CCOL[name]}{row}"] = f"={pid}!$B${b.ok_row(name)}"
    ok, last = CCOL["ok"], SUMMARY_CHECK_FIRST_ROW + len(MODEL_POINTS) - 1
    ws[f"A{SUMMARY_ALL_OK_ROW}"] = "all_ok"
    ws[f"A{SUMMARY_ALL_OK_ROW}"].font = BOLD
    ws[f"B{SUMMARY_ALL_OK_ROW}"] = f"=AND({ok}{SUMMARY_CHECK_FIRST_ROW}:{ok}{last})"
    define(wb, "all_ok", f"{ws.title}!$B${SUMMARY_ALL_OK_ROW}")

    # Section 11.3 totals by t over the batch horizon, each policy's rows beyond its own rem_m
    # read as zero (the sheets have different row counts, so each term is INDEX-guarded).
    horizon = max(mp.rem_m for mp in MODEL_POINTS)
    title(
        ws,
        SUMMARY_TOTALS_T_TITLE_ROW,
        f"Totals by t (section 11.3), t = 0 .. {horizon}: the sum over the ten policies of each"
        " Money and Count monthly column, a policy's values beyond its own rem_m read as zero",
    )
    ws[f"A{SUMMARY_TOTALS_T_HEADER_ROW}"] = "t"
    ws[f"A{SUMMARY_TOTALS_T_HEADER_ROW}"].font = BOLD
    for name, letter in TCOL.items():
        ws[f"{letter}{SUMMARY_TOTALS_T_HEADER_ROW}"] = f"total_{name}"
        ws[f"{letter}{SUMMARY_TOTALS_T_HEADER_ROW}"].font = BOLD
    rem_m_row = DERIVED_FIRST_ROW + MP_DERIVED_FIELDS.index("rem_m")
    for t in range(horizon + 1):
        row = SUMMARY_TOTALS_T_FIRST_ROW + t
        ws[f"A{row}"] = t
        for name, letter in TCOL.items():
            if t == 0 and name not in STATE_TOTALS:
                continue  # month quantities are blank at t = 0 and are not totalled there
            terms: list[str] = []
            for mp in MODEL_POINTS:
                rem = f"{mp.policy_id}!$B${rem_m_row}"
                src = f"{mp.policy_id}!${COL[name]}:${COL[name]}"
                terms.append(f"IF($A{row}<={rem},INDEX({src},$A{row}+{MONTH_FIRST_ROW}),0)")
            ws[f"{letter}{row}"] = "=" + "+".join(terms)
    set_widths(ws, 18, max(len(SUMMARY_SCALAR_COLUMNS), len(TOTAL_COLUMNS) + 1))


def build() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    write_readme(wb)
    write_mortality(wb)
    for mp in MODEL_POINTS:
        write_policy(wb, mp)
    write_summary(wb)
    return wb


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = build()
    wb.save(OUT)
    rows = ", ".join(f"{mp.policy_id}: {mp.rem_m}" for mp in MODEL_POINTS)
    print(f"wrote {OUT} ({len(wb.sheetnames)} sheets; rem_m by policy: {rows})")
    print("formulas only: recalculate with docs/spec/recalc_excel.py before reading values")
    print("then summary!all_ok (defined name all_ok) is the headline check")


if __name__ == "__main__":
    main()

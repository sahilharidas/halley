"""Build the D22 reference workbooks for ``docs/spec/in-force-group.md``, from that spec alone.

The spec defines one reporting period's IFRS 17 measurement of an in-force annuity group in two
variants, P (profitable) and O (onerous). This script writes one workbook per variant,
``docs/spec/in-force-group/in-force-group-P.xlsx`` and ``in-force-group-O.xlsx``, in the layout the
spec's section 11.3 fixes:

* ``inputs``: the run settings of 3.1, the scalar tables of 3.5 with the constants derived from
  them, the group inputs of 3.3 and the cross-run reads of 3.4 (by reference to the REC sheets and
  the group totals), the model points of section 13 with their derived fields, the mortality base
  tables and the improvement table of 6.1, the three curves of 5.1 and 5.3 and the four CPI paths
  of 5.2 and 5.3, all built by formula from the constants the spec prints;
* ``REC-11``, ``OLD-11``, ``NEW-11``, ``CUR-11``, ``REC-12``, ...: one sheet per (run, policy)
  with the columns of 6.10 in evaluation order, one row per month ``t = 0 .. T`` of the run, the
  run's constants read from ``inputs``, the per-policy scalars of 7.2 and the check cells of 12.1;
* ``group``: the totals of 7.1 per run by ``t``, the recognition block of 7.3, the close of
  section 9 one step per row, the months of section 10 by ``t``, the roll-forward tables of 11.1
  and the analysis of change of 11.2;
* ``checks``: one boolean per invariant of section 12 marked (sheet), per (run, policy) or per
  group, and the ``all_ok`` cell (also a workbook-level defined name).

Every quantity the spec defines is an Excel formula transcribing the spec's formula in the spec's
evaluation order (left-to-right chains, ratios of two curve cells, sums in ascending ``t``,
powers rather than running products where the spec says so, ``to_monthly_cf`` as
``1 - (1 - q)^(1/12)`` as 6.1 allows a spreadsheet). The only typed numbers are the reference
inputs of section 13, the constants the spec's formulas name (Makeham parameters, curve and CPI
rates) and the ``t``, ``u``, age and year keys of the tables. The eight (run, policy) sheets share
one layout; they differ only in which ``inputs`` cells and defined names they read (C17). The
script reads nothing but its own constants: not the engine, not the number script, not the
Parquet inputs. openpyxl writes no cached values, so each workbook is recalculated in Excel
afterwards with ``docs/spec/recalc_excel.py`` before anything reads its numbers.

    uv run --frozen python docs/spec/in_force_group_workbook.py
"""

from __future__ import annotations

import datetime as dt
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

OUT_DIR = Path(__file__).resolve().parent / "in-force-group"

# Structural sizes, from the spec's derived-values table of section 13 and the file ranges of
# section 3.6. They fix how many rows the generator writes; the sheets' own T cells are formulas.
T_REC = 501  # the recognition run's horizon
T_V = 489  # the horizon of every run at V
AGES = range(20, 121)  # mort_base and mort_cif age keys
YEARS = range(2025, 2151)  # mort_cif year keys
CPI_REC_U = range(-24, T_REC + 1)  # cpi_rec: u relative to V_rec
CPI_V_U = range(-26, T_V + 1)  # cpi_act, cpi_locked, cpi_cur: u relative to V
EPS = "1E-12"  # section 12: eps within one implementation

BOLD = Font(bold=True)
DATE_FORMAT = "yyyy-mm-dd"


@dataclass(frozen=True)
class ModelPoint:
    """One row of the model point table of section 13, both at recognition and at V."""

    policy_id: int
    klass: str  # "level" or "CPI"
    sex: str
    date_of_birth: dt.date
    pols_if_init_rec: float
    instalment_pm_rec: float
    single_premium: float
    pols_if_init_v: float  # actual lives at V
    instalment_pm_v: float  # actual instalment at V


@dataclass(frozen=True)
class Variant:
    code: str
    profit_bucket: str
    cu_act: float
    policies: tuple[ModelPoint, ModelPoint]


@dataclass(frozen=True)
class Run:
    """One of the four runs of section 5.5: which inputs its (run, policy) sheets read."""

    name: str
    at_recognition: bool
    horizon: int
    curve: str  # defined name of the composed curve df(t)
    cpi: str  # defined name of the CPI path I(u)
    cpi_u: str  # defined name of that path's u axis
    mort_version: int
    maint_pm: str  # inputs key of the run's expense level
    pols_if_init: str  # model point field of the run's opening in-force
    instalment_pm: str  # model point field of the run's instalment


ISSUE_DATE = dt.date(2025, 12, 31)
ANNUITY_START_DATE = dt.date(2026, 1, 1)
VALUATION_DATE = dt.date(2026, 12, 31)

LEVEL_LIFE = (dt.date(1947, 8, 15), "M")
CPI_LIFE = (dt.date(1945, 11, 30), "F")

VARIANTS = (
    Variant(
        "P",
        "OTHER",
        1667000.00,
        (
            ModelPoint(11, "level", "M", LEVEL_LIFE[0], 100, 1000.00, 126250.60, 98, 1000.00),
            ModelPoint(12, "CPI", "F", CPI_LIFE[0], 50, 800.00, 132916.18, 50, 800.00),
        ),
    ),
    Variant(
        "O",
        "ONEROUS",
        1640400.00,
        (
            ModelPoint(21, "level", "M", LEVEL_LIFE[0], 100, 1000.00, 118966.91, 95, 1000.00),
            ModelPoint(22, "CPI", "F", CPI_LIFE[0], 50, 800.00, 125247.94, 47, 800.00),
        ),
    ),
)

RUNS = (
    Run(
        "REC",
        True,
        T_REC,
        "P_rec",
        "I_rec",
        "cpi_rec_u",
        1,
        "maint_pm_v1",
        "pols_if_init_rec",
        "instalment_pm_rec",
    ),
    Run("OLD", False, T_V, "P_L", "I_L", "cpi_u", 1, "maint_pm_old", "n_exp", "instalment_pm_v"),
    Run(
        "NEW",
        False,
        T_V,
        "P_L",
        "I_L",
        "cpi_u",
        2,
        "maint_pm_v2",
        "pols_if_init_v",
        "instalment_pm_v",
    ),
    Run(
        "CUR",
        False,
        T_V,
        "P_cur",
        "I_cur",
        "cpi_u",
        2,
        "maint_pm_v2",
        "pols_if_init_v",
        "instalment_pm_v",
    ),
)


def ref(sheet: str, col: int, row: int) -> str:
    """Absolute reference to one cell of ``sheet``, quoted when the sheet name needs it."""
    name = f"'{sheet}'" if "-" in sheet else sheet
    return f"{name}!${get_column_letter(col)}${row}"


def define(wb: Workbook, name: str, sheet: str, col0: int, row0: int, col1: int, row1: int) -> None:
    """Add a workbook-level defined name for a rectangular range."""
    c0, c1 = get_column_letter(col0), get_column_letter(col1)
    wb.defined_names[name] = DefinedName(name, attr_text=f"{sheet}!${c0}${row0}:${c1}${row1}")


# --------------------------------------------------------------------------------------------
# inputs sheet
# --------------------------------------------------------------------------------------------

# Stacked label/value blocks in columns A:D. Each block is (title, keys); the model point block
# has one value column per policy (B, C) instead of value and note.
INPUT_SETTINGS = (
    "variant",
    "portfolio",
    "cohort",
    "profit_bucket",
    "valuation_date",
    "recognition_date",
    "prev_valuation_date",
    "m_V",
    "m_rec",
    "m_prev",
    "tau0",
    "tau",
    "L",
    "include_inv_exp",
    "discount_coverage_units",
    "J",
    "T_rec",
    "T_V",
)
INPUT_TABLES = (
    "maint_pm_v1",
    "maint_pm_v2",
    "maint_pm_old",
    "ie_rate",
    "ilp",
    "alpha",
    "z_alpha",
    "sigma_l",
    "sigma_e",
    "ie_m",
    "mort_mult",
    "exp_mult",
    "mk_A",
    "mk_B",
    "mk_c",
)
INPUT_GROUP = (
    "g_portfolio",
    "g_cohort",
    "g_profit_bucket",
    "g_recognition_date",
    "g_prev_valuation_date",
    "csm_prev",
    "lc_prev",
    "cu_act",
    "prev_run_id",
    "F_prev0",
    "F_prev_tau",
)
MP_FIELDS = (
    "policy_id",
    "class",
    "sex",
    "date_of_birth",
    "issue_date",
    "annuity_start_date",
    "pols_if_init_rec",
    "instalment_pm_rec",
    "single_premium",
    "pols_if_init_v",
    "instalment_pm_v",
    "portfolio",
    "cohort",
    "profit_bucket",
    "mi_dob",
    "a0_rec",
    "a0_v",
    "mi_asd",
    "m_c_rec",
    "m_c_v",
    "T_i_rec",
    "T_i_v",
    "year_issue",
    "n_exp",
)
INPUT_BLOCKS = (
    ("3.1 Run settings", INPUT_SETTINGS),
    ("3.5 Assumption tables: scalar tables and the constants derived from them", INPUT_TABLES),
    ("3.3 Group inputs and 3.4 reads from the previous run (the REC sheets)", INPUT_GROUP),
    ("13 Model points, one column per policy (derived fields by formula)", MP_FIELDS),
)

INPUT_LABELS = {
    "variant": "variant",
    "portfolio": "portfolio",
    "cohort": "cohort",
    "profit_bucket": "profit_bucket",
    "valuation_date": "valuation_date V",
    "recognition_date": "recognition_date V_rec",
    "prev_valuation_date": "prev_valuation_date V_prev",
    "m_V": "m_V = mi(V)",
    "m_rec": "m_rec = mi(V_rec)",
    "m_prev": "m_prev = mi(V_prev)",
    "tau0": "tau0 = m_V - m_rec",
    "tau": "tau = m_V - m_prev",
    "L": "cpi_lag_m L",
    "include_inv_exp": "include_inv_exp",
    "discount_coverage_units": "discount_coverage_units",
    "J": "scenario count J",
    "T_rec": "T (recognition run) = 12*120 + 1 - min a0 at V_rec",
    "T_V": "T (runs at V) = 12*120 + 1 - min a0 at V",
    "maint_pm_v1": "expense v1: maint_pm (December 2025 prices)",
    "maint_pm_v2": "expense v2: maint_pm (December 2026 prices)",
    "maint_pm_old": "maint_pm_old = maint_pm_v1 * I_act(0) / I_act(-tau)  (5.4)",
    "ie_rate": "inv_expense: ie_rate",
    "ilp": "ilp",
    "alpha": "ra_param: alpha",
    "z_alpha": "ra_param: z_alpha",
    "sigma_l": "ra_param: sigma_l",
    "sigma_e": "ra_param: sigma_e",
    "ie_m": "ie_m = (1 + ie_rate)^(1/12) - 1  (6.4)",
    "mort_mult": "mort_mult = exp(-sigma_l * z_alpha - sigma_l * sigma_l / 2)  (6.8)",
    "exp_mult": "exp_mult = exp(sigma_e * z_alpha - sigma_e * sigma_e / 2)  (6.8)",
    "mk_A": "Makeham A (6.1)",
    "mk_B": "Makeham B (6.1)",
    "mk_c": "Makeham c (6.1)",
    "g_portfolio": "portfolio",
    "g_cohort": "cohort",
    "g_profit_bucket": "profit_bucket",
    "g_recognition_date": "recognition_date",
    "g_prev_valuation_date": "prev_valuation_date",
    "csm_prev": "csm_prev = CSM(0) of the REC run (group sheet)",
    "lc_prev": "lc_prev = LC(0) of the REC run (group sheet)",
    "cu_act": "cu_act (datum: payments actually made in the period)",
    "prev_run_id": "prev_run_id",
    "F_prev0": "F_prev(0) = FCF_g^rec(0)",
    "F_prev_tau": "F_prev(tau) = FCF_g^rec(tau)",
    "policy_id": "policy_id",
    "class": "class",
    "sex": "sex",
    "date_of_birth": "date_of_birth",
    "issue_date": "issue_date",
    "annuity_start_date": "annuity_start_date",
    "pols_if_init_rec": "pols_if_init at recognition",
    "instalment_pm_rec": "instalment_pm at recognition",
    "single_premium": "single_premium",
    "pols_if_init_v": "pols_if_init at V (actual)",
    "instalment_pm_v": "instalment_pm at V (actual)",
    "mi_dob": "mi(date_of_birth)",
    "a0_rec": "a0 at recognition = m_rec - mi(date_of_birth)",
    "a0_v": "a0 at V = m_V - mi(date_of_birth)",
    "mi_asd": "mi(annuity_start_date)",
    "m_c_rec": "m_c at recognition = mi(annuity_start_date) - m_rec",
    "m_c_v": "m_c at V = mi(annuity_start_date) - m_V",
    "T_i_rec": "T_i at recognition",
    "T_i_v": "T_i at V",
    "year_issue": "year(issue_date) (G2: equals cohort)",
    "n_exp": "n_exp = pols_if^rec(tau), read from the REC sheet",
}

# Table columns on the inputs sheet (1-based). Lookup ranges are defined names over these.
MORT_COL = 6  # F: age, G..J: v1 M, v1 F, v2 M, v2 F
CURVE_COL = 12  # L: t, M: s, N: z_rec, O: rf_df_rec, P: P_rec, Q: z_cur, R: rf_df_cur,
#                  S: P_cur, T: P_L, U: d_L, V: j_L
CPI_REC_COL = 24  # X: u (relative to V_rec), Y: I_rec
CPI_V_COL = 27  # AA: u (relative to V), AB: I_act, AC: I_L, AD: I_cur
CIF_COL = 32  # AF: age, AG: g(x), AH..: years; rows 2 and 3 hold h(y) and w(h(y))
CIF_FIRST_ROW = 4


def input_rows() -> tuple[dict[str, int], dict[int, str]]:
    """Row of every inputs key, and the title rows of the blocks."""
    rows: dict[str, int] = {}
    titles: dict[int, str] = {}
    r = 2
    for title, keys in INPUT_BLOCKS:
        titles[r] = title
        r += 1
        for key in keys:
            rows[key] = r
            r += 1
        r += 1
    return rows, titles


INPUT_ROW, INPUT_TITLE = input_rows()


def iref(key: str) -> str:
    """Absolute reference to a scalar input (value column B)."""
    return ref("inputs", 2, INPUT_ROW[key])


def mref(idx: int, key: str) -> str:
    """Absolute reference to a model point field of policy ``idx`` (column B or C)."""
    return ref("inputs", 2 + idx, INPUT_ROW[key])


def write_inputs(wb: Workbook, ws: Worksheet, variant: Variant) -> None:
    ws.column_dimensions["A"].width = 58
    for col in "BCD":
        ws.column_dimensions[col].width = 22
    ws.cell(1, 1, "input").font = BOLD
    ws.cell(1, 2, "value").font = BOLD
    ws.cell(1, 3, "note").font = BOLD
    for r, title in INPUT_TITLE.items():
        ws.cell(r, 1, title).font = BOLD
    for key, label in INPUT_LABELS.items():
        ws.cell(INPUT_ROW[key], 1, label)

    def put(key: str, value: object, note: str = "", fmt: str | None = None) -> None:
        cell = ws.cell(INPUT_ROW[key], 2, value)
        if fmt:
            cell.number_format = fmt
        if note:
            ws.cell(INPUT_ROW[key], 3, note)

    def mi(date_ref: str) -> str:
        return f"12*YEAR({date_ref})+MONTH({date_ref})-1"

    # 3.1 run settings
    put("variant", variant.code)
    put("portfolio", "IA")
    put("cohort", 2025)
    put("profit_bucket", variant.profit_bucket)
    put("valuation_date", VALUATION_DATE, "the last day of a calendar month", DATE_FORMAT)
    put("recognition_date", ISSUE_DATE, "a month-end", DATE_FORMAT)
    put("prev_valuation_date", ISSUE_DATE, "a month-end; here equal to V_rec (C1)", DATE_FORMAT)
    put("m_V", f"={mi(iref('valuation_date'))}", "mi(d) = 12 * year(d) + month(d) - 1")
    put("m_rec", f"={mi(iref('recognition_date'))}")
    put("m_prev", f"={mi(iref('prev_valuation_date'))}")
    put("tau0", f"={iref('m_V')}-{iref('m_rec')}", "months from recognition to V")
    put("tau", f"={iref('m_V')}-{iref('m_prev')}", "the reporting period in months")
    put("L", 3, "build-time constant of the CPI class")
    put("include_inv_exp", True)
    put("discount_coverage_units", True, "the only form in the spec (C13)")
    put("J", 1, "one scenario per run (G6)")
    put(
        "T_rec",
        f"=12*120+1-MIN({mref(0, 'a0_rec')},{mref(1, 'a0_rec')})",
        f"the generator writes rows t = 0 .. {T_REC} on the REC sheets",
    )
    put(
        "T_V",
        f"=12*120+1-MIN({mref(0, 'a0_v')},{mref(1, 'a0_v')})",
        f"the generator writes rows t = 0 .. {T_V} on the OLD, NEW and CUR sheets",
    )

    # 3.5 scalar tables and derived constants
    put("maint_pm_v1", 5.00, "version 1, priced at V_prev; used by REC as is, by OLD rebased")
    put("maint_pm_v2", 5.30, "version 2, priced at V; used by NEW and CUR")
    i_act0 = "INDEX(I_act,MATCH(0,cpi_u,0))"
    i_act_tau = f"INDEX(I_act,MATCH(-{iref('tau')},cpi_u,0))"
    put("maint_pm_old", f"={iref('maint_pm_v1')}*{i_act0}/{i_act_tau}", "used by OLD (5.4, C8)")
    put("ie_rate", 0.0015, "Rate[Annual, Effective]; not stressed")
    put("ilp", 0.0050, "Rate[Annual, Continuous]; part of the locked-in basis")
    put("alpha", 0.75)
    put("z_alpha", 0.6744897501960817)
    put("sigma_l", 0.10)
    put("sigma_e", 0.10)
    put("ie_m", f"=(1+{iref('ie_rate')})^(1/12)-1")
    sl, se, za = iref("sigma_l"), iref("sigma_e"), iref("z_alpha")
    put("mort_mult", f"=EXP(-{sl}*{za}-{sl}*{sl}/2)")
    put("exp_mult", f"=EXP({se}*{za}-{se}*{se}/2)")
    put("mk_A", 0.00022)
    put("mk_B", 2.7e-6)
    put("mk_c", 1.124)

    # 3.3 group inputs and 3.4 reads from the previous run
    put("g_portfolio", f"={iref('portfolio')}")
    put("g_cohort", f"={iref('cohort')}")
    put("g_profit_bucket", f"={iref('profit_bucket')}")
    put("g_recognition_date", f"={iref('recognition_date')}", "", DATE_FORMAT)
    put("g_prev_valuation_date", f"={iref('prev_valuation_date')}", "", DATE_FORMAT)
    put("csm_prev", f"={gref('CSM0_rec')}", "read from the recognition run, not retyped (3.4)")
    put("lc_prev", f"={gref('LC0_rec')}", "read from the recognition run, not retyped (3.4)")
    put("cu_act", variant.cu_act, "section 13: the group input, a total")
    put("prev_run_id", f"REC-{variant.code}", "the recognition run of this variant: the REC sheets")
    g_t = f"group!${GC['t']}$2:${GC['t']}${T_REC + 2}"
    g_fcf_rec = f"group!${GC['FCF_rec']}$2:${GC['FCF_rec']}${T_REC + 2}"
    put("F_prev0", f"=INDEX({g_fcf_rec},MATCH(0,{g_t},0))", "FCF_g^rec at month 0 (7.1)")
    put(
        "F_prev_tau",
        f"=INDEX({g_fcf_rec},MATCH({iref('tau')},{g_t},0))",
        "FCF_g^rec at month tau (7.1)",
    )

    # 13 model points, one column per policy
    hdr_row = INPUT_ROW["policy_id"] - 1
    for idx, mp in enumerate(variant.policies):
        ws.cell(hdr_row, 2 + idx, f"policy {mp.policy_id}").font = BOLD
        write_model_point(ws, variant, idx, mp)
    ws.cell(INPUT_ROW["n_exp"], 4, "3.4: the previous run's pols_if at its month tau")
    ws.cell(INPUT_ROW["pols_if_init_v"], 4, "the survivors of the period (3.2)")
    ws.cell(INPUT_ROW["instalment_pm_v"], 4, "equals the recognition instalment here (section 4)")

    write_mortality_tables(wb, ws)
    write_curves(wb, ws)
    write_cpi_paths(wb, ws)
    write_cif_table(wb, ws)


def write_model_point(ws: Worksheet, variant: Variant, idx: int, mp: ModelPoint) -> None:
    """One model point of section 13 as a column: the typed inputs, then the derived fields."""
    col = 2 + idx

    def mput(key: str, value: object, fmt: str | None = None) -> None:
        cell = ws.cell(INPUT_ROW[key], col, value)
        if fmt:
            cell.number_format = fmt

    def me(key: str) -> str:
        return mref(idx, key)

    def mi(date_ref: str) -> str:
        return f"12*YEAR({date_ref})+MONTH({date_ref})-1"

    rec_sheet = f"REC-{mp.policy_id}"
    rec_t = f"'{rec_sheet}'!${PC['t']}$2:${PC['t']}${T_REC + 2}"
    rec_pols = f"'{rec_sheet}'!${PC['pols_if']}$2:${PC['pols_if']}${T_REC + 2}"
    mput("policy_id", mp.policy_id)
    mput("class", mp.klass)
    mput("sex", mp.sex)
    mput("date_of_birth", mp.date_of_birth, DATE_FORMAT)
    mput("issue_date", ISSUE_DATE, DATE_FORMAT)
    mput("annuity_start_date", ANNUITY_START_DATE, DATE_FORMAT)
    mput("pols_if_init_rec", mp.pols_if_init_rec)
    mput("instalment_pm_rec", mp.instalment_pm_rec)
    mput("single_premium", mp.single_premium)
    mput("pols_if_init_v", mp.pols_if_init_v)
    mput("instalment_pm_v", mp.instalment_pm_v)
    mput("portfolio", "IA")
    mput("cohort", 2025)
    mput("profit_bucket", variant.profit_bucket)
    mput("mi_dob", f"={mi(me('date_of_birth'))}")
    mput("a0_rec", f"={iref('m_rec')}-{me('mi_dob')}")
    mput("a0_v", f"={iref('m_V')}-{me('mi_dob')}")
    mput("mi_asd", f"={mi(me('annuity_start_date'))}")
    mput("m_c_rec", f"={me('mi_asd')}-{iref('m_rec')}")
    mput("m_c_v", f"={me('mi_asd')}-{iref('m_V')}")
    mput("T_i_rec", f"=12*120+1-{me('a0_rec')}")
    mput("T_i_v", f"=12*120+1-{me('a0_v')}")
    mput("year_issue", f"=YEAR({me('issue_date')})")
    mput("n_exp", f"=INDEX({rec_pols},MATCH({iref('tau')},{rec_t},0))")


def write_mortality_tables(wb: Workbook, ws: Worksheet) -> None:
    """mort_base v1 (Makeham, 6.1) and v2 (0.96 x v1, 1 at 120), one row per age."""
    c = MORT_COL
    heads = (
        "age",
        "mort_base_v1(M, x)",
        "mort_base_v1(F, x)",
        "mort_base_v2(M, x)",
        "mort_base_v2(F, x)",
    )
    for i, h in enumerate(heads):
        ws.cell(1, c + i, h).font = BOLD
        ws.column_dimensions[get_column_letter(c + i)].width = 20
    a, b, k = iref("mk_A"), iref("mk_B"), iref("mk_c")

    def q_mk(x: str) -> str:
        # evaluate: c^x, then B * c^x, then * (c - 1), then / ln(c), then -A - that, exp, 1 -
        return f"1-EXP(-{a}-{b}*{k}^{x}*({k}-1)/LN({k}))"

    for i, age in enumerate(AGES):
        r = 2 + i
        age_ref = f"${get_column_letter(c)}{r}"
        ws.cell(r, c, age)
        if age == 120:
            for j in range(1, 5):
                ws.cell(r, c + j, "=1")
        else:
            ws.cell(r, c + 1, f"={q_mk(age_ref)}")
            ws.cell(r, c + 2, f"={q_mk(f'({age_ref}-4)')}")
            ws.cell(r, c + 3, f"=0.96*{get_column_letter(c + 1)}{r}")
            ws.cell(r, c + 4, f"=0.96*{get_column_letter(c + 2)}{r}")
    last = 1 + len(AGES)
    define(wb, "mort_age", "inputs", c, 2, c, last)
    define(wb, "mort_v1_M", "inputs", c + 1, 2, c + 1, last)
    define(wb, "mort_v1_F", "inputs", c + 2, 2, c + 2, last)
    define(wb, "mort_v2_M", "inputs", c + 3, 2, c + 3, last)
    define(wb, "mort_v2_F", "inputs", c + 4, 2, c + 4, last)


def write_curves(wb: Workbook, ws: Worksheet) -> None:
    """The composed curves of 5.1 and 5.3 and the locked-in curve re-based to V, by t."""
    c = CURVE_COL
    heads = (
        "t",
        "s = t / 12",
        "z_rec(s)",
        "rf_df_rec(t)",
        "P_rec(t)",
        "z_cur(s)",
        "rf_df_cur(t)",
        "P_cur(t)",
        "P_L(t)",
        "d_L(t)",
        "j_L(t)",
    )
    for i, h in enumerate(heads):
        ws.cell(1, c + i, h).font = BOLD
        ws.column_dimensions[get_column_letter(c + i)].width = 18
    col = {h.split("(")[0].split(" ")[0]: get_column_letter(c + i) for i, h in enumerate(heads)}
    ilp, tau0 = iref("ilp"), iref("tau0")
    p_rec_at_tau0 = f"INDEX(P_rec,MATCH({tau0},curve_t,0))"
    for t in range(T_REC + 1):
        r = 2 + t
        t_, s = f"{col['t']}{r}", f"{col['s']}{r}"
        ws.cell(r, c, t)
        ws.cell(r, c + 1, f"={t_}/12")
        ws.cell(r, c + 2, f"=0.040-0.015*EXP(-{s}/10)")
        ws.cell(r, c + 3, f"=EXP(-{col['z_rec']}{r}*{s})")
        ws.cell(r, c + 4, f"={col['rf_df_rec']}{r}*EXP(-{ilp}*{t_}/12)")
        if t <= T_V:
            ws.cell(r, c + 5, f"=0.045-0.015*EXP(-{s}/10)")
            ws.cell(r, c + 6, f"=EXP(-{col['z_cur']}{r}*{s})")
            ws.cell(r, c + 7, f"={col['rf_df_cur']}{r}*EXP(-{ilp}*{t_}/12)")
            ws.cell(r, c + 8, f"=INDEX(P_rec,MATCH({tau0}+{t_},curve_t,0))/{p_rec_at_tau0}")
            if t >= 1:
                p_l, p_l_prev = f"{col['P_L']}{r}", f"{col['P_L']}{r - 1}"
                ws.cell(r, c + 9, f"={p_l}/{p_l_prev}")
                ws.cell(r, c + 10, f"={p_l_prev}/{p_l}-1")
    last = 2 + T_REC
    define(wb, "curve_t", "inputs", c, 2, c, last)
    define(wb, "P_rec", "inputs", c + 4, 2, c + 4, last)
    define(wb, "P_cur", "inputs", c + 7, 2, c + 7, last)
    define(wb, "P_L", "inputs", c + 8, 2, c + 8, last)
    define(wb, "d_L", "inputs", c + 9, 2, c + 9, last)
    define(wb, "j_L", "inputs", c + 10, 2, c + 10, last)


def write_cpi_paths(wb: Workbook, ws: Worksheet) -> None:
    """I_rec by u relative to V_rec; I_act, I_L and I_cur by u relative to V (5.2, 5.3)."""
    c = CPI_REC_COL
    ws.cell(1, c, "u (relative to V_rec)").font = BOLD
    ws.cell(1, c + 1, "I_rec(u)").font = BOLD
    for i, u in enumerate(CPI_REC_U):
        r = 2 + i
        u_ = f"{get_column_letter(c)}{r}"
        ws.cell(r, c, u)
        ws.cell(r, c + 1, f"=IF({u_}<=0,100*1.04^({u_}/12),100*1.025^({u_}/12))")
    last = 1 + len(CPI_REC_U)
    define(wb, "cpi_rec_u", "inputs", c, 2, c, last)
    define(wb, "I_rec", "inputs", c + 1, 2, c + 1, last)

    c = CPI_V_COL
    for i, h in enumerate(("u (relative to V)", "I_act(u)", "I_L(u)", "I_cur(u)")):
        ws.cell(1, c + i, h).font = BOLD
        ws.column_dimensions[get_column_letter(c + i)].width = 18
    tau0 = iref("tau0")
    u_col, act_col = get_column_letter(c), get_column_letter(c + 1)
    i_act0 = f"${act_col}${2 + CPI_V_U.index(0)}"
    i_rec_at_tau0 = f"INDEX(I_rec,MATCH({tau0},cpi_rec_u,0))"
    for i, u in enumerate(CPI_V_U):
        r = 2 + i
        u_ = f"{u_col}{r}"
        ws.cell(r, c, u)
        if u <= 0:
            ws.cell(r, c + 1, f"=IF({u_}<=-13,100*1.04^(({u_}+12)/12),100.1*1.032^(({u_}+12)/12))")
            ws.cell(r, c + 2, f"={act_col}{r}")
            ws.cell(r, c + 3, f"={act_col}{r}")
        else:
            ws.cell(
                r, c + 2, f"={i_act0}*INDEX(I_rec,MATCH({tau0}+{u_},cpi_rec_u,0))/{i_rec_at_tau0}"
            )
            ws.cell(r, c + 3, f"={i_act0}*1.03^({u_}/12)")
    last = 1 + len(CPI_V_U)
    define(wb, "cpi_u", "inputs", c, 2, c, last)
    define(wb, "I_act", "inputs", c + 1, 2, c + 1, last)
    define(wb, "I_L", "inputs", c + 2, 2, c + 2, last)
    define(wb, "I_cur", "inputs", c + 3, 2, c + 3, last)


def write_cif_table(wb: Workbook, ws: Worksheet) -> None:
    """mort_cif of 6.1: cif(x, 2025) = 1, cif(x, y) = cif(x, y - 1) * (1 - r(x, y)).

    Row 1 holds the year keys, rows 2 and 3 the year's h(y) and w(h(y)), column AF the age keys
    and column AG the age's g(x); r(x, y) = g(x) * (0.010 + 0.010 * w(h(y))) is inlined in each
    cif cell.
    """
    c = CIF_COL
    ws.cell(1, c, "age (rows) by year (columns)").font = BOLD
    ws.cell(1, c + 1, "g(x)").font = BOLD
    ws.cell(2, c, "h(y) = min(1, (y - 2026) / 20)")
    ws.cell(3, c, "w(h) = 1 - 3 h^2 + 2 h^3")
    g_col = get_column_letter(c + 1)
    for j, year in enumerate(YEARS):
        yc = get_column_letter(c + 2 + j)
        ws.cell(1, c + 2 + j, year).font = BOLD
        if year >= 2026:
            ws.cell(2, c + 2 + j, f"=MIN(1,({yc}$1-2026)/20)")
            ws.cell(3, c + 2 + j, f"=1-3*{yc}$2^2+2*{yc}$2^3")
    for i, age in enumerate(AGES):
        r = CIF_FIRST_ROW + i
        age_ref = f"${get_column_letter(c)}{r}"
        ws.cell(r, c, age)
        ws.cell(r, c + 1, f"=IF({age_ref}<=85,1,IF({age_ref}<110,(110-{age_ref})/25,0))")
        for j, year in enumerate(YEARS):
            yc = get_column_letter(c + 2 + j)
            if year == 2025:
                ws.cell(r, c + 2 + j, "=1")
            else:
                prev = get_column_letter(c + 1 + j)
                ws.cell(r, c + 2 + j, f"={prev}{r}*(1-${g_col}{r}*(0.010+0.010*{yc}$3))")
    last_row = CIF_FIRST_ROW + len(AGES) - 1
    last_col = c + 1 + len(YEARS)
    define(wb, "cif_age", "inputs", c, CIF_FIRST_ROW, c, last_row)
    define(wb, "cif_year", "inputs", c + 2, 1, last_col, 1)
    define(wb, "cif", "inputs", c + 2, CIF_FIRST_ROW, last_col, last_row)


# --------------------------------------------------------------------------------------------
# (run, policy) sheets
# --------------------------------------------------------------------------------------------

# Columns in the evaluation order of 6.10, then the helper columns (7.2 terms, I9's n_K) and the
# per-row booleans of the 12.1 checks. Header row 1 names the spec's symbols.
POLICY_COLS = (
    "t",
    "x",
    "Y",
    "q_ann",
    "q_mth",
    "deaths",
    "pols_if",
    "anniv",
    "esc_factor",
    "ann_pp",
    "exp_pp",
    "ann_out",
    "exp_out",
    "df",
    "v",
    "s_pp",
    "q_ann_st",
    "q_mth_st",
    "s_pp_st",
    "cu",
    "bel",
    "inv_exp",
    "v_be",
    "v_st",
    "ra",
    "pva_term",
    "pve_term",
    "n_K",
    "I9_rhs",
    "chk_I1",
    "chk_I2",
    "chk_I3",
    "chk_I4",
    "chk_I7",
    "chk_I9",
)
PC = {name: get_column_letter(i + 1) for i, name in enumerate(POLICY_COLS)}
POLICY_HEADERS = {
    "pva_term": "df(t-1) * ann_out(t)",
    "pve_term": "df(t-1) * exp_out(t) + df(t) * inv_exp(t)",
    "n_K": "n_K (I9)",
    "I9_rhs": "instalment_pm * I(n_K - L) / I(n_0 - L)",
}
POLICY_CONSTS = (
    "run",
    "policy_id",
    "class",
    "sex",
    "pols_if_init",
    "instalment_pm",
    "m_V",
    "a0",
    "m_c",
    "T_i",
    "T",
    "L",
    "ie_m",
    "mort_mult",
    "exp_mult",
    "maint_pm",
    "I0",
    "mort_base",
    "curve",
    "cpi_path",
)
CONST_LABELS = {
    "I0": "I(0)",
    "m_V": "m_V (the run's valuation month)",
    "mort_base": "mort_base (defined name read)",
    "curve": "df(t) (defined name read)",
    "cpi_path": "I(u) (defined name read)",
}
POLICY_SCALARS = ("bel0", "ra0", "pva", "pve", "ann_pp1", "v_be0", "v_st0")
SCALAR_LABELS = {
    "bel0": "bel(0)",
    "ra0": "ra(0)",
    "pva": "pva",
    "pve": "pve",
    "ann_pp1": "ann_pp(1)",
    "v_be0": "v_be(0)",
    "v_st0": "v_st(0)",
}
POLICY_CHECKS = ("I1", "I2", "I3", "I4", "I5", "I7", "I9")

KCOL = len(POLICY_COLS) + 2  # label column of the constants block; value one to the right
SCOL = KCOL + 3  # label column of the scalars block
CCOL = SCOL + 3  # label column of the check block


def kaddr(key: str) -> str:
    return f"${get_column_letter(KCOL + 1)}${2 + POLICY_CONSTS.index(key)}"


def saddr(key: str) -> str:
    return f"${get_column_letter(SCOL + 1)}${2 + POLICY_SCALARS.index(key)}"


def caddr(key: str) -> str:
    return f"${get_column_letter(CCOL + 1)}${2 + POLICY_CHECKS.index(key)}"


def write_policy_sheet(wb: Workbook, run: Run, idx: int, mp: ModelPoint) -> None:
    ws = wb.create_sheet(f"{run.name}-{mp.policy_id}")
    horizon = run.horizon
    is_cpi = mp.klass == "CPI"
    for i, name in enumerate(POLICY_COLS):
        ws.cell(1, i + 1, POLICY_HEADERS.get(name, name)).font = BOLD
    ws.freeze_panes = "B2"

    def cell(name: str, t: int) -> str:
        return f"{PC[name]}{t + 2}"

    def col_range(name: str, t0: int, t1: int) -> str:
        return f"{PC[name]}{t0 + 2}:{PC[name]}{t1 + 2}"

    # The run's and policy's constants (6.10 step 1), read from the inputs sheet.
    ws.cell(1, KCOL, "constant").font = BOLD
    ws.cell(1, KCOL + 1, "value").font = BOLD
    ws.column_dimensions[get_column_letter(KCOL)].width = 30
    rec = run.at_recognition
    consts: dict[str, object] = {
        "run": run.name,
        "policy_id": f"={mref(idx, 'policy_id')}",
        "class": f"={mref(idx, 'class')}",
        "sex": f"={mref(idx, 'sex')}",
        "pols_if_init": f"={mref(idx, run.pols_if_init)}",
        "instalment_pm": f"={mref(idx, run.instalment_pm)}",
        "m_V": f"={iref('m_rec' if rec else 'm_V')}",
        "a0": f"={mref(idx, 'a0_rec' if rec else 'a0_v')}",
        "m_c": f"={mref(idx, 'm_c_rec' if rec else 'm_c_v')}",
        "T_i": f"={mref(idx, 'T_i_rec' if rec else 'T_i_v')}",
        "T": f"={iref('T_rec' if rec else 'T_V')}",
        "L": f"={iref('L')}",
        "ie_m": f"={iref('ie_m')}",
        "mort_mult": f"={iref('mort_mult')}",
        "exp_mult": f"={iref('exp_mult')}",
        "maint_pm": f"={iref(run.maint_pm)}",
        "I0": f"=INDEX({run.cpi},MATCH(0,{run.cpi_u},0))",
        "mort_base": f"mort_v{run.mort_version}_{mp.sex}",
        "curve": run.curve,
        "cpi_path": run.cpi,
    }
    for i, key in enumerate(POLICY_CONSTS):
        ws.cell(2 + i, KCOL, CONST_LABELS.get(key, key))
        ws.cell(2 + i, KCOL + 1, consts[key])
    k = {key: kaddr(key) for key in POLICY_CONSTS}
    mort = f"mort_v{run.mort_version}_{mp.sex}"

    def cpi_at(u_expr: str) -> str:
        return f"INDEX({run.cpi},MATCH({u_expr},{run.cpi_u},0))"

    # Row t = 0 (6.10 step 2), then rows t = 1 .. T left to right (step 3), then the backward
    # columns (step 4). Every formula reads its own row and row t - 1 (or t + 1 for the backward
    # columns) only.
    for t in range(horizon + 1):
        r = t + 2
        ws.cell(r, 1, t)
        t_ = cell("t", t)
        if t >= 1:
            x, y = cell("x", t), cell("Y", t)
            ws.cell(r, 2, f"=MIN(120,INT(({k['a0']}+{t_}-1)/12))")
            ws.cell(r, 3, f"=INT(({k['m_V']}+{t_})/12)")
            ws.cell(
                r,
                4,
                f"=IF({x}>=120,1,INDEX({mort},MATCH({x},mort_age,0))"
                f"*INDEX(cif,MATCH({x},cif_age,0),MATCH({y},cif_year,0)))",
            )
            ws.cell(r, 5, f"=1-(1-{cell('q_ann', t)})^(1/12)")
            ws.cell(r, 6, f"={cell('pols_if', t - 1)}*{cell('q_mth', t)}")
            ws.cell(r, 7, f"={cell('pols_if', t - 1)}-{cell('deaths', t)}")
            ws.cell(r, 8, f"=AND({t_}-{k['m_c']}>0,MOD({t_}-{k['m_c']},12)=0)")
            i_lag = cpi_at(f"{t_}-{k['L']}")
            i_lag_12 = cpi_at(f"{t_}-{k['L']}-12")
            esc = f"IF({cell('anniv', t)},{i_lag}/{i_lag_12},1)"
            ws.cell(r, 9, f'=IF({k["class"]}="CPI",{esc},1)')
            ws.cell(r, 10, f"={cell('ann_pp', t - 1)}*{cell('esc_factor', t)}")
            i_prev = cpi_at(f"{t_}-1")
            ws.cell(r, 11, f"={k['maint_pm']}*({i_prev}/{k['I0']})")
            ws.cell(r, 12, f"={cell('pols_if', t - 1)}*{cell('ann_pp', t)}")
            ws.cell(r, 13, f"={cell('pols_if', t - 1)}*{cell('exp_pp', t)}")
        else:
            ws.cell(r, 7, f"={k['pols_if_init']}")
            ws.cell(r, 10, f"={k['instalment_pm']}")
        ws.cell(r, 14, f"=INDEX({run.curve},MATCH({t_},curve_t,0))")
        if t >= 1:
            ws.cell(r, 15, f"={cell('df', t)}/{cell('df', t - 1)}")
            ws.cell(r, 16, f"={cell('ann_pp', t)}+{cell('exp_pp', t)}")
            ws.cell(r, 17, f"=IF({cell('x', t)}>=120,1,{cell('q_ann', t)}*{k['mort_mult']})")
            ws.cell(r, 18, f"=1-(1-{cell('q_ann_st', t)})^(1/12)")
            ws.cell(r, 19, f"={cell('ann_pp', t)}+{cell('exp_pp', t)}*{k['exp_mult']}")
            ws.cell(r, 20, f"={cell('ann_out', t)}")
        # backward columns: bel, inv_exp, v_be, v_st (6.7, 6.8), then ra for every row
        if t == horizon:
            ws.cell(r, 21, "=0")
            ws.cell(r, 23, "=0")
            ws.cell(r, 24, "=0")
        else:
            n = t + 1
            v_n, ie = cell("v", n), k["ie_m"]
            ws.cell(
                r,
                21,
                f"={cell('ann_out', n)}+{cell('exp_out', n)}"
                f"+{cell('inv_exp', n)}*{v_n}+{v_n}*{cell('bel', n)}",
            )
            ws.cell(
                r,
                23,
                f"={cell('s_pp', n)}+(1-{cell('q_mth', n)})"
                f"*({v_n}*({cell('v_be', n)}+{ie}*{cell('v_be', n)}))",
            )
            ws.cell(
                r,
                24,
                f"={cell('s_pp_st', n)}+(1-{cell('q_mth_st', n)})"
                f"*({v_n}*({cell('v_st', n)}+{ie}*{cell('v_st', n)}))",
            )
        if t >= 1:
            ws.cell(r, 22, f"={k['ie_m']}*{cell('bel', t)}")
        ws.cell(r, 25, f"={cell('pols_if', t)}*({cell('v_st', t)}-{cell('v_be', t)})")
        # 7.2 terms, summed in ascending t by the scalars block
        if t >= 1:
            ws.cell(r, 26, f"={cell('df', t - 1)}*{cell('ann_out', t)}")
            ws.cell(
                r,
                27,
                f"={cell('df', t - 1)}*{cell('exp_out', t)}+{cell('df', t)}*{cell('inv_exp', t)}",
            )
        # I9 helpers (CPI class): n_0 in row 0, then the latest anniversary month in 1 .. t
        if is_cpi:
            if t == 0:
                ws.cell(r, 28, f"=IF({k['m_c']}=1,{k['m_c']},{k['m_c']}+12*INT(-{k['m_c']}/12))")
            else:
                ws.cell(r, 28, f"=IF({cell('anniv', t)},{t_},{cell('n_K', t - 1)})")
                i_n_k = cpi_at(f"{cell('n_K', t)}-{k['L']}")
                i_n_0 = cpi_at(f"{cell('n_K', 0)}-{k['L']}")
                ws.cell(r, 29, f"={k['instalment_pm']}*{i_n_k}/{i_n_0}")
        # per-row booleans of the 12.1 checks
        if t >= 1:
            pf, pf_prev = cell("pols_if", t), cell("pols_if", t - 1)
            ws.cell(r, 30, f"=AND({pf}>=0,{pf}<={pf_prev})")
            qs = [cell(q, t) for q in ("q_ann", "q_mth", "q_ann_st", "q_mth_st")]
            bounds = ",".join(f"{q}>=0,{q}<=1" for q in qs)
            ws.cell(r, 31, f"=AND({bounds},{cell('deaths', t)}<={pf_prev})")
            bel, bel_prev = cell("bel", t), cell("bel", t - 1)
            ao, eo = cell("ann_out", t), cell("exp_out", t)
            v, ie_t = cell("v", t), cell("inv_exp", t)
            ws.cell(
                r,
                33,
                f"=ABS({bel}-(({bel_prev}-{ao}-{eo})/{v}-{ie_t}))"
                f"<={EPS}*((ABS({bel_prev})+ABS({ao})+ABS({eo}))/{v}+ABS({ie_t}))",
            )
            if is_cpi:
                rhs = cell("I9_rhs", t)
                ws.cell(r, 35, f"=ABS({cell('ann_pp', t)}-{rhs})<={EPS}*ABS({rhs})")
        ws.cell(
            r,
            32,
            f"=IF({t_}>{k['T_i']},AND({cell('pols_if', t)}=0,{cell('bel', t)}=0,"
            f"{cell('ra', t)}=0),TRUE)",
        )
        ws.cell(r, 34, f"={cell('ra', t)}>=0")

    # 7.2 per-policy scalars (6.10 step 5)
    ws.cell(1, SCOL, "scalar (7.2)").font = BOLD
    ws.cell(1, SCOL + 1, "value").font = BOLD
    ws.column_dimensions[get_column_letter(SCOL)].width = 16
    scalars = {
        "bel0": f"={cell('bel', 0)}",
        "ra0": f"={cell('ra', 0)}",
        "pva": f"=SUM({col_range('pva_term', 1, horizon)})",
        "pve": f"=SUM({col_range('pve_term', 1, horizon)})",
        "ann_pp1": f"={cell('ann_pp', 1)}",
        "v_be0": f"={cell('v_be', 0)}",
        "v_st0": f"={cell('v_st', 0)}",
    }
    for i, key in enumerate(POLICY_SCALARS):
        ws.cell(2 + i, SCOL, SCALAR_LABELS[key])
        ws.cell(2 + i, SCOL + 1, scalars[key])

    # 12.1 check cells, one boolean per invariant for this (run, policy)
    ws.cell(1, CCOL, "check (12.1)").font = BOLD
    ws.cell(1, CCOL + 1, "value").font = BOLD
    ws.column_dimensions[get_column_letter(CCOL)].width = 16
    t_axis = col_range("t", 0, horizon)
    at_ti = f"MATCH({k['T_i']},{t_axis},0)"
    pva_pve = f"({saddr('pva')}+{saddr('pve')})"
    checks: dict[str, object] = {
        "I1": f"=AND({col_range('chk_I1', 1, horizon)})",
        "I2": f"=AND({col_range('chk_I2', 1, horizon)})",
        "I3": f"=AND({cell('bel', horizon)}=0,{cell('v_be', horizon)}=0,{cell('v_st', horizon)}=0,"
        f"INDEX({col_range('q_ann', 0, horizon)},{at_ti})=1,"
        f"INDEX({col_range('pols_if', 0, horizon)},{at_ti})=0,"
        f"AND({col_range('chk_I3', 0, horizon)}))",
        "I4": f"=AND({col_range('chk_I4', 1, horizon)})",
        "I5": f"=ABS({saddr('bel0')}-{pva_pve})<={EPS}*MAX(ABS({saddr('bel0')}),ABS({pva_pve}))",
        "I7": f"=AND({col_range('chk_I7', 0, horizon)})",
        "I9": f"=AND({col_range('chk_I9', 1, horizon)})" if is_cpi else "n/a (level class)",
    }
    for i, key in enumerate(POLICY_CHECKS):
        ws.cell(2 + i, CCOL, key)
        ws.cell(2 + i, CCOL + 1, checks[key])


# --------------------------------------------------------------------------------------------
# group sheet
# --------------------------------------------------------------------------------------------

GROUP_COLS = (
    "t",
    "BEL_rec",
    "RA_rec",
    "FCF_rec",
    "CU_rec",
    "BEL_old",
    "RA_old",
    "FCF_old",
    "CU_old",
    "BEL_new",
    "RA_new",
    "FCF_new",
    "CU_new",
    "BEL_cur",
    "RA_cur",
    "FCF_cur",
    "CU_cur",
    "CU_g",
    "F",
    "d_L",
    "j_L",
    "N",
    "Acc",
    "CSM_pre",
    "U",
    "R",
    "CSM",
    "LC",
    "LCA",
    "chk_K6",
    "chk_I11",
    "chk_I12",
    "chk_I13",
)
GC = {name: get_column_letter(i + 1) for i, name in enumerate(GROUP_COLS)}
RUN_SUFFIX = {"REC": "rec", "OLD": "old", "NEW": "new", "CUR": "cur"}
RUN_SYMBOL = {"REC": "rec", "OLD": "L_old", "NEW": "L_new", "CUR": "cur"}


def group_headers() -> dict[str, str]:
    heads = {
        "t": "t",
        "CU_g": "CU_g(t) (CUR run, section 10)",
        "F": "F(t) = FCF_g^cur(t)",
        "d_L": "d_L(t)",
        "j_L": "j_L(t)",
        "N": "N(t)",
        "Acc": "Acc(t)",
        "CSM_pre": "CSM_pre(t)",
        "U": "U(t)",
        "R": "R(t)",
        "CSM": "CSM(t)",
        "LC": "LC(t)",
        "LCA": "LCA(t)",
        "chk_K6": "chk_K6",
        "chk_I11": "chk_I11",
        "chk_I12": "chk_I12",
        "chk_I13": "chk_I13",
    }
    for run, sym in RUN_SYMBOL.items():
        sfx = RUN_SUFFIX[run]
        heads[f"BEL_{sfx}"] = f"BEL_g^{sym}(t)"
        heads[f"RA_{sfx}"] = f"RA_g^{sym}(t)"
        heads[f"FCF_{sfx}"] = f"FCF_g^{sym}(t)"
        heads[f"CU_{sfx}"] = f"CU_g^{sym}(t)"
    return heads


# The scalar blocks of the group sheet, one quantity per row: (key, label). A key starting with
# "hdr_" is a title row.
GROUP_SCALARS: tuple[tuple[str, str], ...] = (
    ("hdr_rec", "7.3 Initial recognition (from the REC run)"),
    ("P_g", "P_g = sum_i pols_if_init_rec_i * single_premium_i"),
    ("FCF_rec0", "FCF_g^rec(0)"),
    ("gap", "gap = P_g - FCF_g^rec(0)"),
    ("CSM0_rec", "CSM(0) = where(gap > 0, gap, 0)"),
    ("LC0_rec", "LC(0) = max(0, FCF_g^rec(0) - P_g)"),
    ("G3_1", "G3 diagnostic: first policy onerous at recognition"),
    ("G3_2", "G3 diagnostic: second policy onerous at recognition"),
    ("CU_exp_period", "diagnostic: expected units of the period, sum CU_g^rec(1 .. tau)"),
    ("hdr_close", "9 The period close, one step per row"),
    ("csm_prev", "csm_prev"),
    ("lc_prev", "lc_prev"),
    ("cu_act", "cu_act"),
    ("tau0", "tau0"),
    ("tau", "tau"),
    ("P_rec_a", "P_rec(tau0 - tau)"),
    ("P_rec_b", "P_rec(tau0)"),
    ("CSM_acc", "step 1: CSM_acc = csm_prev * P_rec(tau0 - tau) / P_rec(tau0)"),
    ("acc_per", "acc_per = CSM_acc - csm_prev"),
    ("F_prev0", "F_prev(0)"),
    ("F_prev_tau", "F_prev(tau)"),
    ("LC_alloc", "step 2: LC_alloc = where(F_prev(0) > 0, lc_prev * (F_prev(tau) / F_prev(0)), 0)"),
    ("lca_per", "lca_per = lc_prev - LC_alloc"),
    ("FCF_new0", "FCF_g^L_new(0)"),
    ("FCF_old0", "FCF_g^L_old(0)"),
    ("delta_FCF_L", "section 8: delta_FCF_L = FCF_g^L_new(0) - FCF_g^L_old(0)"),
    ("rule", "step 3: 9.4 rule applied (1: LC > 0, delta >= 0; 2: LC > 0, delta < 0; 3: LC = 0)"),
    ("rev", "rev = min(LC_alloc, -delta_FCF_L)  (rule 2)"),
    ("CSM_adj_3", "CSM_acc - delta_FCF_L  (rule 3, before the para 48(a) floor)"),
    ("LC_adj", "LC_adj"),
    ("CSM_adj", "CSM_adj"),
    ("csm_fs", "csm_fs = CSM_adj - CSM_acc"),
    ("lc_fs", "lc_fs = LC_adj - LC_alloc"),
    ("N0", "step 4: N(0) (months block, from the CUR run's units on P_L)"),
    ("R_per", "R_per = where(cu_act + N(0) > 0, CSM_adj * (cu_act / (cu_act + N(0))), CSM_adj)"),
    ("CSM0", "step 5: CSM(0) = CSM_adj - R_per"),
    ("LC0", "LC(0) = LC_adj"),
    ("FCF_cur0", "FCF_g^cur(0)"),
    ("LRC0", "LRC(0) = FCF_g^cur(0) + CSM(0)"),
    ("hdr_rf_csm", "11.1 CSM roll-forward for the period"),
    ("rf_csm_open", "Opening CSM"),
    ("rf_csm_acc", "Interest accretion (44(b), B72(b))"),
    ("rf_csm_fs", "Changes for future service (44(c), B96)"),
    ("rf_csm_fx", "Currency differences (44(d))"),
    ("rf_csm_rel", "Release (44(e), B119)"),
    ("rf_csm_close", "Closing CSM"),
    ("hdr_rf_lc", "11.1 Loss component roll-forward for the period"),
    ("rf_lc_open", "Opening LC"),
    ("rf_lc_alloc", "Systematic allocation (50(a), 51, 52)"),
    ("rf_lc_loss", "Losses on remeasurement (48(a)): lc_fs where positive"),
    ("rf_lc_rev", "Reversals of losses (50(b)): lc_fs where negative"),
    ("rf_lc_close", "Closing LC"),
    ("hdr_aoc", "11.2 Analysis of change of the fulfilment cash flows"),
    ("aoc_A", "A = FCF_g^rec(0)"),
    ("aoc_B", "B = F_prev(tau)"),
    ("aoc_C", "C = FCF_g^L_old(0)"),
    ("aoc_D", "D = FCF_g^L_new(0)"),
    ("aoc_E", "E = FCF_g^cur(0)"),
    ("aoc_BA", "B - A: the previous run's expected movement"),
    ("aoc_CB", "C - B: CPI experience (IFIE)"),
    ("aoc_DC", "D - C: delta_FCF_L (CSM or LC)"),
    ("aoc_ED", "E - D: financial assumptions at V (IFIE)"),
    ("hdr_helpers", "Helpers (layout, not definition)"),
    ("t_last", "last month with CU_g > 0 (I11)"),
)
GROUP_SCALAR_ROW = {key: 2 + i for i, (key, _) in enumerate(GROUP_SCALARS)}
GLCOL = len(GROUP_COLS) + 2  # label column of the scalar blocks; value one to the right


def gref(key: str) -> str:
    """Absolute reference to a group-sheet scalar."""
    return ref("group", GLCOL + 1, GROUP_SCALAR_ROW[key])


def gcol_range(name: str, t0: int, t1: int) -> str:
    return f"{GC[name]}{t0 + 2}:{GC[name]}{t1 + 2}"


def write_group(wb: Workbook, variant: Variant) -> None:
    ws = wb.create_sheet("group")
    heads = group_headers()
    for i, name in enumerate(GROUP_COLS):
        ws.cell(1, i + 1, heads[name]).font = BOLD
    ws.freeze_panes = "B2"
    pids = [mp.policy_id for mp in variant.policies]  # ascending policy_id (D19)

    def gcell(name: str, t: int) -> str:
        return f"{GC[name]}{t + 2}"

    def total(run: str, col: str, t: int) -> str:
        return "+".join(f"'{run}-{pid}'!{PC[col]}{t + 2}" for pid in pids)

    # 7.1 totals per run by t
    for run in RUNS:
        sfx = RUN_SUFFIX[run.name]
        for t in range(run.horizon + 1):
            r = t + 2
            ws.cell(r, 1, t)
            ws.cell(r, GROUP_COLS.index(f"BEL_{sfx}") + 1, f"={total(run.name, 'bel', t)}")
            ws.cell(r, GROUP_COLS.index(f"RA_{sfx}") + 1, f"={total(run.name, 'ra', t)}")
            ws.cell(
                r,
                GROUP_COLS.index(f"FCF_{sfx}") + 1,
                f"={gcell(f'BEL_{sfx}', t)}+{gcell(f'RA_{sfx}', t)}",
            )
            if t >= 1:
                ws.cell(r, GROUP_COLS.index(f"CU_{sfx}") + 1, f"={total(run.name, 'cu', t)}")

    # Section 10: months after the close, t = 0 .. T, on the CUR run's units and F
    csm0, lc0 = gref("CSM0"), gref("LC0")
    for t in range(T_V + 1):
        r = t + 2
        ws.cell(r, GROUP_COLS.index("F") + 1, f"={gcell('FCF_cur', t)}")
        if t >= 1:
            ws.cell(r, GROUP_COLS.index("CU_g") + 1, f"={gcell('CU_cur', t)}")
            ws.cell(r, GROUP_COLS.index("d_L") + 1, f"=INDEX(d_L,MATCH({gcell('t', t)},curve_t,0))")
            ws.cell(r, GROUP_COLS.index("j_L") + 1, f"=INDEX(j_L,MATCH({gcell('t', t)},curve_t,0))")
        if t == T_V:
            ws.cell(r, GROUP_COLS.index("N") + 1, "=0")
        else:
            ws.cell(
                r,
                GROUP_COLS.index("N") + 1,
                f"={gcell('CU_g', t + 1)}+{gcell('d_L', t + 1)}*{gcell('N', t + 1)}",
            )
        if t == 0:
            ws.cell(r, GROUP_COLS.index("CSM") + 1, f"={csm0}")
            ws.cell(r, GROUP_COLS.index("LC") + 1, f"={lc0}")
            continue
        cu, n, u = gcell("CU_g", t), gcell("N", t), gcell("U", t)
        csm, csm_prev, acc, pre, rel = (
            gcell("CSM", t),
            gcell("CSM", t - 1),
            gcell("Acc", t),
            gcell("CSM_pre", t),
            gcell("R", t),
        )
        lc, lc_prev, f, f_prev = (
            gcell("LC", t),
            gcell("LC", t - 1),
            gcell("F", t),
            gcell("F", t - 1),
        )
        ws.cell(r, GROUP_COLS.index("Acc") + 1, f"={gcell('j_L', t)}*{csm_prev}")
        ws.cell(r, GROUP_COLS.index("CSM_pre") + 1, f"={csm_prev}+{acc}")
        ws.cell(r, GROUP_COLS.index("U") + 1, f"={cu}+{n}")
        ws.cell(r, GROUP_COLS.index("R") + 1, f"=IF({u}>0,{pre}*({cu}/{u}),{pre})")
        ws.cell(r, GROUP_COLS.index("CSM") + 1, f"={pre}-{rel}")
        ws.cell(r, GROUP_COLS.index("LC") + 1, f"=IF({f_prev}>0,{lc_prev}*({f}/{f_prev}),0)")
        ws.cell(r, GROUP_COLS.index("LCA") + 1, f"={lc_prev}-{lc}")
        # per-row booleans of the 12.2 checks
        ws.cell(r, GROUP_COLS.index("chk_K6") + 1, f"=AND({csm}>=0,{lc}>=0)")
        ws.cell(
            r,
            GROUP_COLS.index("chk_I11") + 1,
            f"=AND(ABS({csm}-({csm_prev}+{acc}-{rel}))<={EPS}*({csm_prev}+ABS({acc})+{rel}),"
            f"IF({gcell('t', t)}>{gref('t_last')},{csm}=0,TRUE))",
        )
        ws.cell(
            r,
            GROUP_COLS.index("chk_I12") + 1,
            f"=IF(AND({cu}>0,{n}>0,{csm}>0),ABS({rel}/{cu}-{csm}/{n})"
            f"<={EPS}*MAX(ABS({rel}/{cu}),ABS({csm}/{n})),TRUE)",
        )
        f0 = gcell("F", 0)
        ws.cell(
            r,
            GROUP_COLS.index("chk_I13") + 1,
            f"=IF({f}>0,IF({lc0}>0,ABS({lc}/{f}-{lc0}/{f0})"
            f"<={EPS}*MAX(ABS({lc}/{f}),ABS({lc0}/{f0})),TRUE),{lc}=0)",
        )

    # Scalar blocks: recognition, close, roll-forward tables, analysis of change
    ws.cell(1, GLCOL, "quantity").font = BOLD
    ws.cell(1, GLCOL + 1, "value").font = BOLD
    ws.column_dimensions[get_column_letter(GLCOL)].width = 80
    ws.column_dimensions[get_column_letter(GLCOL + 1)].width = 22
    for key, label in GROUP_SCALARS:
        cell = ws.cell(GROUP_SCALAR_ROW[key], GLCOL, label)
        if key.startswith("hdr_"):
            cell.font = BOLD
    g = {key: gref(key) for key, _ in GROUP_SCALARS}
    premium_terms = "+".join(
        f"{mref(i, 'pols_if_init_rec')}*{mref(i, 'single_premium')}" for i in range(len(pids))
    )
    t_axis = gcol_range("t", 0, T_REC)
    cu_rec = gcol_range("CU_rec", 0, T_REC)
    delta = g["delta_FCF_L"]
    values: dict[str, object] = {
        "P_g": f"={premium_terms}",
        "FCF_rec0": f"={gcell('FCF_rec', 0)}",
        "gap": f"={g['P_g']}-{g['FCF_rec0']}",
        "CSM0_rec": f"=IF({g['gap']}>0,{g['gap']},0)",
        "LC0_rec": f"=MAX(0,{g['FCF_rec0']}-{g['P_g']})",
        "CU_exp_period": f"=SUMPRODUCT(({t_axis}>=1)*({t_axis}<={iref('tau')})*{cu_rec})",
        "csm_prev": f"={iref('csm_prev')}",
        "lc_prev": f"={iref('lc_prev')}",
        "cu_act": f"={iref('cu_act')}",
        "tau0": f"={iref('tau0')}",
        "tau": f"={iref('tau')}",
        "P_rec_a": f"=INDEX(P_rec,MATCH({g['tau0']}-{g['tau']},curve_t,0))",
        "P_rec_b": f"=INDEX(P_rec,MATCH({g['tau0']},curve_t,0))",
        "CSM_acc": f"={g['csm_prev']}*{g['P_rec_a']}/{g['P_rec_b']}",
        "acc_per": f"={g['CSM_acc']}-{g['csm_prev']}",
        "F_prev0": f"={iref('F_prev0')}",
        "F_prev_tau": f"={iref('F_prev_tau')}",
        "LC_alloc": f"=IF({g['F_prev0']}>0,{g['lc_prev']}*({g['F_prev_tau']}/{g['F_prev0']}),0)",
        "lca_per": f"={g['lc_prev']}-{g['LC_alloc']}",
        "FCF_new0": f"={gcell('FCF_new', 0)}",
        "FCF_old0": f"={gcell('FCF_old', 0)}",
        "delta_FCF_L": f"={g['FCF_new0']}-{g['FCF_old0']}",
        "rule": f"=IF(AND({g['LC_alloc']}>0,{delta}>=0),1,"
        f"IF(AND({g['LC_alloc']}>0,{delta}<0),2,3))",
        "rev": f"=MIN({g['LC_alloc']},-{delta})",
        "CSM_adj_3": f"={g['CSM_acc']}-{delta}",
        "LC_adj": f"=IF({g['rule']}=1,{g['LC_alloc']}+{delta},IF({g['rule']}=2,"
        f"{g['LC_alloc']}-{g['rev']},IF({g['CSM_adj_3']}<0,-{g['CSM_adj_3']},0)))",
        "CSM_adj": f"=IF({g['rule']}=1,{g['CSM_acc']},IF({g['rule']}=2,"
        f"{g['CSM_acc']}+(-{delta}-{g['rev']}),IF({g['CSM_adj_3']}<0,0,{g['CSM_adj_3']})))",
        "csm_fs": f"={g['CSM_adj']}-{g['CSM_acc']}",
        "lc_fs": f"={g['LC_adj']}-{g['LC_alloc']}",
        "N0": f"={gcell('N', 0)}",
        "R_per": f"=IF({g['cu_act']}+{g['N0']}>0,{g['CSM_adj']}*({g['cu_act']}"
        f"/({g['cu_act']}+{g['N0']})),{g['CSM_adj']})",
        "CSM0": f"={g['CSM_adj']}-{g['R_per']}",
        "LC0": f"={g['LC_adj']}",
        "FCF_cur0": f"={gcell('FCF_cur', 0)}",
        "LRC0": f"={g['FCF_cur0']}+{g['CSM0']}",
        "rf_csm_open": f"={g['csm_prev']}",
        "rf_csm_acc": f"={g['acc_per']}",
        "rf_csm_fs": f"={g['csm_fs']}",
        "rf_csm_fx": 0,
        "rf_csm_rel": f"=-{g['R_per']}",
        "rf_csm_close": f"={g['CSM0']}",
        "rf_lc_open": f"={g['lc_prev']}",
        "rf_lc_alloc": f"=-{g['lca_per']}",
        "rf_lc_loss": f"=IF({g['lc_fs']}>0,{g['lc_fs']},0)",
        "rf_lc_rev": f"=IF({g['lc_fs']}<0,{g['lc_fs']},0)",
        "rf_lc_close": f"={g['LC0']}",
        "aoc_A": f"={g['FCF_rec0']}",
        "aoc_B": f"={g['F_prev_tau']}",
        "aoc_C": f"={g['FCF_old0']}",
        "aoc_D": f"={g['FCF_new0']}",
        "aoc_E": f"={g['FCF_cur0']}",
        "aoc_BA": f"={g['aoc_B']}-{g['aoc_A']}",
        "aoc_CB": f"={g['aoc_C']}-{g['aoc_B']}",
        "aoc_DC": f"={g['aoc_D']}-{g['aoc_C']}",
        "aoc_ED": f"={g['aoc_E']}-{g['aoc_D']}",
        "t_last": f"=SUMPRODUCT(MAX(({gcol_range('CU_g', 1, T_V)}>0)*{gcol_range('t', 1, T_V)}))",
    }
    for i, pid in enumerate(pids):
        rec = f"REC-{pid}"
        fcf_i = f"('{rec}'!{saddr('bel0')}+'{rec}'!{saddr('ra0')})"
        values[f"G3_{i + 1}"] = (
            f"={mref(i, 'pols_if_init_rec')}*{mref(i, 'single_premium')}-{fcf_i}<0"
        )
    for key, value in values.items():
        ws.cell(GROUP_SCALAR_ROW[key], GLCOL + 1, value)


# --------------------------------------------------------------------------------------------
# checks sheet
# --------------------------------------------------------------------------------------------

GROUP_CHECKS = ("K1", "K2", "K3", "K4", "K5", "K6", "I11", "I12", "I13")


def write_checks(wb: Workbook, variant: Variant) -> None:
    ws = wb.create_sheet("checks")
    ws.column_dimensions["A"].width = 34
    ws.cell(1, 1, "(run, policy) sheet").font = BOLD
    for j, key in enumerate(POLICY_CHECKS):
        ws.cell(1, 2 + j, key).font = BOLD
    r = 2
    for mp in variant.policies:
        for run in RUNS:
            sheet = f"{run.name}-{mp.policy_id}"
            ws.cell(r, 1, sheet)
            for j, key in enumerate(POLICY_CHECKS):
                if key == "I9" and mp.klass != "CPI":
                    ws.cell(r, 2 + j, "n/a (level class)")
                else:
                    ws.cell(r, 2 + j, f"='{sheet}'!{caddr(key)}")
            r += 1
    policy_block = f"B2:{get_column_letter(1 + len(POLICY_CHECKS))}{r - 1}"

    r += 1
    ws.cell(r, 1, "group check (12.2)").font = BOLD
    ws.cell(r, 2, "value").font = BOLD
    r += 1
    first_group = r
    g = {key: gref(key) for key, _ in GROUP_SCALARS}
    p_g = g["P_g"]
    group_values = {
        "K1": f"=ABS({g['CSM0']}-({g['csm_prev']}+{g['acc_per']}+{g['csm_fs']}-{g['R_per']}))"
        f"<={EPS}*{p_g}",
        "K2": f"=ABS({g['LC0']}-({g['lc_prev']}-{g['lca_per']}+{g['lc_fs']}))<={EPS}*{p_g}",
        "K3": f"=ABS(({g['CSM_adj']}-{g['LC_adj']})-(({g['CSM_acc']}-{g['LC_alloc']})"
        f"-{g['delta_FCF_L']}))<={EPS}*{p_g}",
        "K4": f"=AND({g['CSM_acc']}>=0,{g['LC_alloc']}>=0,{g['CSM_adj']}>=0,{g['LC_adj']}>=0,"
        f"{g['CSM0']}>=0,{g['LC0']}>=0,NOT(AND({g['CSM_adj']}>0,{g['LC_adj']}>0)),"
        f"NOT(AND({g['CSM0']}>0,{g['LC0']}>0)))",
        "K5": f"=IF(AND({g['cu_act']}>0,{g['CSM0']}>0),ABS({g['R_per']}/{g['cu_act']}"
        f"-{g['CSM0']}/{g['N0']})<={EPS}*MAX(ABS({g['R_per']}/{g['cu_act']}),"
        f"ABS({g['CSM0']}/{g['N0']})),TRUE)",
        "K6": f"=AND(group!{gcol_range('chk_K6', 1, T_V)})",
        "I11": f"=AND(group!{gcol_range('chk_I11', 1, T_V)})",
        "I12": f"=AND(group!{gcol_range('chk_I12', 1, T_V)})",
        "I13": f"=AND(group!{gcol_range('chk_I13', 1, T_V)})",
    }
    for key in GROUP_CHECKS:
        ws.cell(r, 1, key)
        ws.cell(r, 2, group_values[key])
        r += 1
    group_block = f"B{first_group}:B{r - 1}"

    r += 1
    ws.cell(r, 1, "all_ok").font = BOLD
    ws.cell(r, 2, f"=AND({policy_block},{group_block})")
    wb.defined_names["all_ok"] = DefinedName("all_ok", attr_text=f"checks!$B${r}")


# --------------------------------------------------------------------------------------------
# assembly
# --------------------------------------------------------------------------------------------


def build(variant: Variant) -> Path:
    wb = Workbook()
    inputs = wb.active
    inputs.title = "inputs"
    write_inputs(wb, inputs, variant)
    for mp_idx, mp in enumerate(variant.policies):
        for run in RUNS:
            write_policy_sheet(wb, run, mp_idx, mp)
    write_group(wb, variant)
    write_checks(wb, variant)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"in-force-group-{variant.code}.xlsx"
    wb.save(path)
    return path


def main() -> int:
    for variant in VARIANTS:
        path = build(variant)
        print(f"wrote {path}")
    print(
        "recalculate with: uv run --frozen python docs/spec/recalc_excel.py "
        "docs/spec/in-force-group/in-force-group-P.xlsx "
        "docs/spec/in-force-group/in-force-group-O.xlsx"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
